"""[CD1.8b] Kiểm scripts/chu_thich_chan_doan.py: đi vòng CSV -> .xlsx -> CSV không mất, không lệch."""

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import chu_thich_chan_doan as m  # noqa: E402
from build_diagnostic_candidates import COT, NHOM  # noqa: E402


def _dong(cau: int, kc: str, nhom: str, nhan: str = "negative", text: str = "") -> dict:
    r = dict.fromkeys(COT, "")
    r.update(uid=f"visfd-test-{cau:05d}-{kc}", text=text or f"câu số {cau}, có dấu phẩy", aspect=kc,
             nhan_goc=nhan, nhom_ung_vien=nhom, tu_khoa_kich_hoat="không", so_token="6")
    return r


def _csv(tmp_path: Path) -> Path:
    rows = []
    for i, g in enumerate(NHOM):
        for j in range(3):
            cau = i * 10 + j
            rows.append(_dong(cau, "BATTERY", g))
            rows.append(_dong(cau, "SCREEN", g, "positive"))
    rows.append(_dong(99, "PRICE", NHOM[0], "neutral", text="=1+1 không phải công thức"))
    p = tmp_path / "candidates.csv"
    m.ghi_csv(p, rows)
    return p


def _dien(xlsx: Path, gia_tri: dict) -> None:
    """gia_tri: {uid: (phenomenon, nhan_kiem_tra, ghi_chu)}"""
    wb = load_workbook(xlsx)
    ws = wb[m.TEN_SHEET]
    cot = {c: i for i, c in enumerate(m.COT_XLSX)}
    for hang in ws.iter_rows(min_row=2):
        uid = hang[cot["uid"]].value
        if uid in gia_tri:
            for c, v in zip(m.COT_CHU_THICH, gia_tri[uid]):
                hang[cot[c]].value = v
    wb.save(xlsx)


def test_thu_tu_xen_ke_bon_nhom_va_giu_khia_canh_cua_mot_cau_lien_nhau(tmp_path):
    rows = m.doc_csv(_csv(tmp_path))
    tt = m.thu_tu_chu_thich(rows, seed=42)
    assert sorted(r["uid"] for r in tt) == sorted(r["uid"] for r in rows)
    assert [r["uid"] for r in tt] == [r["uid"] for r in m.thu_tu_chu_thich(rows, seed=42)]
    cau = [m.ma_cau(r["uid"]) for r in tt]
    gon = [c for i, c in enumerate(cau) if i == 0 or cau[i - 1] != c]
    assert len(gon) == len(set(cau))  # mỗi câu chỉ xuất hiện thành một khối liền
    nhom_dau = {r["nhom_ung_vien"] for r in tt[:8]}  # 4 câu đầu, mỗi câu 2 khía cạnh
    assert nhom_dau == set(NHOM)


def test_di_vong_khong_mat_gi_va_giu_thu_tu_csv(tmp_path):
    p = _csv(tmp_path)
    truoc = m.doc_csv(p)
    x = tmp_path / "a.xlsx"
    assert m.main(["excel", "--csv", str(p), "--xlsx", str(x)]) == 0
    u1, u2 = truoc[0]["uid"], truoc[-1]["uid"]
    _dien(x, {u1: ("negation", "dong_y", "ghi chú, có phẩy"), u2: ("regular", "positive", "")})
    assert m.main(["tu-excel", "--xlsx", str(x), "--csv", str(p)]) == 0
    sau = m.doc_csv(p)
    assert [r["uid"] for r in sau] == [r["uid"] for r in truoc]
    assert (sau[0]["phenomenon"], sau[0]["nhan_kiem_tra"], sau[0]["ghi_chu"]) == ("negation", "dong_y", "ghi chú, có phẩy")
    assert sau[-1]["text"] == "=1+1 không phải công thức" and sau[-1]["nhan_kiem_tra"] == "positive"
    for a, b in zip(truoc, sau):
        assert {k: v for k, v in a.items() if k not in m.COT_CHU_THICH} == \
               {k: v for k, v in b.items() if k not in m.COT_CHU_THICH}
    assert len(list(tmp_path.glob("candidates.*.csv.bak"))) == 1
    assert p.read_bytes().startswith(b"\xef\xbb\xbf") and b"\r\n" in p.read_bytes()


def test_gia_tri_sai_hoac_dien_thieu_thi_khong_ghi(tmp_path):
    p = _csv(tmp_path)
    x = tmp_path / "a.xlsx"
    m.main(["excel", "--csv", str(p), "--xlsx", str(x)])
    truoc = p.read_bytes()
    uid = m.doc_csv(p)[0]["uid"]
    _dien(x, {uid: ("phu dinh", "dong_y", "")})
    assert m.main(["tu-excel", "--xlsx", str(x), "--csv", str(p)]) == 1
    _dien(x, {uid: ("negation", None, "")})
    assert m.main(["tu-excel", "--xlsx", str(x), "--csv", str(p)]) == 1
    assert p.read_bytes() == truoc


def test_sua_cot_du_lieu_trong_excel_thi_bi_chan(tmp_path):
    p = _csv(tmp_path)
    x = tmp_path / "a.xlsx"
    m.main(["excel", "--csv", str(p), "--xlsx", str(x)])
    wb = load_workbook(x)
    wb[m.TEN_SHEET].cell(row=2, column=m.COT_XLSX.index("nhan_goc") + 1).value = "positive?"
    wb.save(x)
    truoc = p.read_bytes()
    assert m.main(["tu-excel", "--xlsx", str(x), "--csv", str(p)]) == 1
    assert p.read_bytes() == truoc


def test_excel_khong_de_len_file_dang_co_va_mang_theo_chu_thich_cu(tmp_path):
    p = _csv(tmp_path)
    x = tmp_path / "a.xlsx"
    m.main(["excel", "--csv", str(p), "--xlsx", str(x)])
    uid = m.doc_csv(p)[3]["uid"]
    _dien(x, {uid: ("contrast", "khong_ro", "x")})
    m.main(["tu-excel", "--xlsx", str(x), "--csv", str(p)])
    assert m.main(["excel", "--csv", str(p), "--xlsx", str(x)]) == 2
    assert m.main(["excel", "--csv", str(p), "--xlsx", str(x), "--ghi-de"]) == 0
    assert len(list(tmp_path.glob("a.*.xlsx.bak"))) == 1
    ws = load_workbook(x)[m.TEN_SHEET]
    hang = next(h for h in ws.iter_rows(min_row=2, values_only=True) if h[1] == uid)
    assert hang[m.COT_XLSX.index("phenomenon")] == "contrast"


def test_chi_ba_cot_chu_thich_mo_khoa_va_co_o_chon(tmp_path):
    p = _csv(tmp_path)
    x = tmp_path / "a.xlsx"
    m.main(["excel", "--csv", str(p), "--xlsx", str(x)])
    ws = load_workbook(x)[m.TEN_SHEET]
    assert ws.protection.sheet
    mo = {c for c, o in zip(m.COT_XLSX, ws[2]) if not o.protection.locked}
    assert mo == set(m.COT_CHU_THICH)
    cong_thuc = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert f'"{",".join(m.HIEN_TUONG)}"' in cong_thuc and f'"{",".join(m.NHAN_KIEM_TRA)}"' in cong_thuc


def test_tien_do_dem_theo_hien_tuong(tmp_path, capsys):
    p = _csv(tmp_path)
    rows = m.doc_csv(p)
    rows[0].update(phenomenon="negation", nhan_kiem_tra="dong_y")
    rows[1].update(phenomenon="regular", nhan_kiem_tra="negative")
    m.ghi_csv(p, rows)
    assert m.main(["tien-do", "--csv", str(p)]) == 0
    ra = capsys.readouterr().out
    assert f"Đã gán 2 / {len(rows)} dòng." in ra and "dong_y=1" in ra
    dong = next(d for d in ra.splitlines() if d.startswith("negation"))
    assert dong.split()[-1] == "99"  # còn thiếu 99 để đủ 100
