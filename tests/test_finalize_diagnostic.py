"""[CD1.8d] Kiểm scripts/finalize_diagnostic.py: rút mẫu gán lại, tính kappa, chấm theo nhóm và chốt tập."""

import datetime as dt
import json
import sys
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import chu_thich_chan_doan as ct  # noqa: E402
import finalize_diagnostic as m  # noqa: E402
from build_diagnostic_candidates import COT  # noqa: E402


def _csv(tmp_path: Path, moi_nhom: int = 6) -> Path:
    rows = []
    for i, ph in enumerate(m.NHOM_RUT):
        for j in range(moi_nhom):
            r = dict.fromkeys(COT, "")
            r.update(uid=f"visfd-test-{i * 100 + j:05d}-BATTERY", text=f"câu {i}-{j}", aspect="BATTERY", nhan_goc="negative",
                     nhom_ung_vien="phu_dinh", tu_khoa_kich_hoat="không", so_token="5", phenomenon=ph, nhan_kiem_tra="dong_y")
            rows.append(r)
    p = tmp_path / "candidates.csv"
    ct.ghi_csv(p, rows)
    return p


def _gan_lai(xlsx: Path, lan1: dict, doi: int = 0) -> None:
    """Điền lần hai giống lần một, trừ `doi` dòng đầu bị đổi hiện tượng."""
    wb = load_workbook(xlsx)
    ws = wb[ct.TEN_SHEET]
    c = {k: i for i, k in enumerate(ct.COT_XLSX)}
    for n, hang in enumerate(ws.iter_rows(min_row=2)):
        ph = lan1[hang[c["uid"]].value]["phenomenon"]
        if n < doi:
            ph = "regular" if ph != "regular" else "negation"
        hang[c["phenomenon"]].value = ph
        hang[c["nhan_kiem_tra"]].value = "dong_y"
    wb.save(xlsx)


def test_kappa_khop_sklearn_va_hai_ca_bien():
    a = ["negation", "contrast", "regular", "regular", "both", "negation", "regular", "contrast"]
    b = ["negation", "regular", "regular", "regular", "both", "contrast", "regular", "contrast"]
    k, p_o, _ = m.kappa_cohen(a, b)
    assert k == pytest.approx(cohen_kappa_score(a, b)) and p_o == pytest.approx(6 / 8)
    assert m.kappa_cohen(a, a)[0] == pytest.approx(1.0)
    assert m.kappa_cohen(["x", "x"], ["x", "x"])[0] == 1.0
    with pytest.raises(ValueError):
        m.kappa_cohen(["x"], [])


def test_rut_phan_tang_deu_tat_dinh_va_bao_loi_khi_thieu(tmp_path):
    rows = ct.doc_csv(_csv(tmp_path))
    chon = m.rut_phan_tang(rows, 4, seed=1)
    assert len(chon) == 16 and len({r["uid"] for r in chon}) == 16
    assert {ph: sum(r["phenomenon"] == ph for r in chon) for ph in m.NHOM_RUT} == dict.fromkeys(m.NHOM_RUT, 4)
    assert [r["uid"] for r in chon] == [r["uid"] for r in m.rut_phan_tang(rows, 4, seed=1)]
    with pytest.raises(ValueError):
        m.rut_phan_tang(rows, 7, seed=1)


def test_file_gan_lai_khong_lo_nhan_cu_va_khong_bi_de(tmp_path):
    p, x = _csv(tmp_path), tmp_path / "lai.xlsx"
    assert m.main(["rut-mau", "--csv", str(p), "--xlsx", str(x), "--moi-nhom", "3"]) == 0
    ws = load_workbook(x)[ct.TEN_SHEET]
    c = {k: i for i, k in enumerate(ct.COT_XLSX)}
    for hang in ws.iter_rows(min_row=2, values_only=True):
        assert all(hang[c[k]] in (None, "") for k in ct.COT_CHU_THICH)
    assert "nhom_ung_vien" not in ct.COT_XLSX
    so = json.loads((tmp_path / "lai.json").read_text(encoding="utf-8"))
    assert len(so["uid"]) == 12
    assert m.main(["rut-mau", "--csv", str(p), "--xlsx", str(x), "--moi-nhom", "3"]) == 2


def test_kappa_chan_khi_chua_du_14_ngay_va_khi_gan_chua_xong(tmp_path):
    p, x = _csv(tmp_path), tmp_path / "lai.xlsx"
    m.main(["rut-mau", "--csv", str(p), "--xlsx", str(x), "--moi-nhom", "3"])
    out = tmp_path / "k.json"
    assert m.main(["kappa", "--csv", str(p), "--xlsx", str(x), "--out", str(out)]) == 2
    assert m.main(["kappa", "--csv", str(p), "--xlsx", str(x), "--out", str(out), "--bo-qua-ngay"]) == 1  # chưa gán
    assert not out.exists()


def test_kappa_dat_va_khong_dat(tmp_path, capsys):
    p, x = _csv(tmp_path), tmp_path / "lai.xlsx"
    m.main(["rut-mau", "--csv", str(p), "--xlsx", str(x), "--moi-nhom", "5"])
    lan1 = {r["uid"]: r for r in ct.doc_csv(p)}
    so = tmp_path / "lai.json"
    d = json.loads(so.read_text(encoding="utf-8"))
    d["som_nhat_duoc_gan_lai"] = (dt.date.today() - dt.timedelta(days=1)).strftime(m.DINH_DANG_NGAY)
    so.write_text(json.dumps(d), encoding="utf-8")
    out = tmp_path / "k.json"
    _gan_lai(x, lan1, doi=1)
    assert m.main(["kappa", "--csv", str(p), "--xlsx", str(x), "--out", str(out)]) == 0
    kq = json.loads(out.read_text(encoding="utf-8"))
    assert kq["loai"] == "nhat_quan_mot_nguoi" and kq["so_dong"] == 20 and len(kq["uid_lech"]) == 1
    assert kq["phenomenon"]["ti_le_trung"] == pytest.approx(19 / 20) and kq["phenomenon"]["kappa"] > 0.9
    _gan_lai(x, lan1, doi=10)
    assert m.main(["kappa", "--csv", str(p), "--xlsx", str(x), "--out", str(out)]) == 3
    assert "CHƯA ĐẠT" in capsys.readouterr().out


# --- Lệnh `cham` và `chot` -------------------------------------------------------


def _du_doan(goc: Path, rows, sai=()) -> None:
    """Ghi predictions.jsonl cho mọi mô hình, mọi seed: đoán đúng hết, trừ các uid trong `sai`."""
    for exp in m.MO_HINH:
        for s in m.cb.SEEDS:
            d = goc / exp / f"seed{s}"
            d.mkdir(parents=True)
            with (d / "predictions.jsonl").open("w", encoding="utf-8") as fh:
                for r in rows:
                    g = m.MA_NHAN[r["nhan_goc"]]
                    fh.write(json.dumps({"uid": r["uid"], "y_true": g, "y_pred": 2 if r["uid"] in sai else g}) + "\n")


def _kappa_file(p: Path, k: float) -> Path:
    p.write_text(json.dumps({"phenomenon": {"kappa": k}}), encoding="utf-8")
    return p


def test_cham_dem_dung_theo_nhom_va_hieu_voi_nhom_thuong(tmp_path):
    rows = ct.doc_csv(_csv(tmp_path, moi_nhom=10))
    sai = {r["uid"] for r in rows if r["phenomenon"] == "both"} | {[r["uid"] for r in rows if r["phenomenon"] == "contrast"][0]}
    _du_doan(tmp_path / "kq", rows, sai)
    cau, o = m.dem_theo_nhom(rows, tmp_path / "kq", ["phobert"])
    assert len(cau) == 40 and o["phobert"].shape == (4, len(m.cb.SEEDS), 40, 9) and o["phobert"].sum() == 40 * len(m.cb.SEEDS)
    r = m.cham_mot_mo_hinh(o["phobert"], 300, np.random.default_rng(0))
    assert {g: r[g]["n"] for g in m.NHOM_CHAM} == dict.fromkeys(m.NHOM_CHAM, 10)
    assert r["regular"]["accuracy"]["gia_tri"] == 1.0 and "p" not in r["regular"]["accuracy"]
    assert r["contrast"]["accuracy"]["gia_tri"] == pytest.approx(0.9)
    assert r["both"]["accuracy"]["hieu_voi_moc"] == pytest.approx(-1.0) and r["both"]["accuracy"]["p"] < 0.05
    assert r["negation"]["accuracy"]["p"] == 1.0  # không khác nhóm thường ở lần rút nào
    x = r["contrast"]["accuracy"]
    assert x["ci95_duoi"] <= x["gia_tri"] <= x["ci95_tren"]


def test_cham_dung_khi_nhan_that_lech_hoac_thieu_du_doan(tmp_path):
    rows = ct.doc_csv(_csv(tmp_path))
    _du_doan(tmp_path / "kq", rows)
    lech = [dict(r) for r in rows]
    lech[0]["nhan_goc"] = "positive"
    with pytest.raises(ValueError, match="khác nhan_goc"):
        m.dem_theo_nhom(lech, tmp_path / "kq", ["phobert"])
    them = rows + [{**rows[0], "uid": "visfd-test-99999-BATTERY"}]
    with pytest.raises(ValueError, match="không có dự đoán"):
        m.dem_theo_nhom(them, tmp_path / "kq", ["phobert"])


def test_cham_ghi_so_bo_cho_toi_khi_kappa_dat(tmp_path, capsys):
    p = _csv(tmp_path)
    rows = ct.doc_csv(p)
    rows[0]["phenomenon"] = "bo"          # dòng bỏ không được chấm
    rows[1]["nhan_kiem_tra"] = "positive"  # dòng lệch nhãn gốc chỉ nằm trong bảng "toàn bộ"
    ct.ghi_csv(p, rows)
    _du_doan(tmp_path / "kq", rows)
    chung = ["cham", "--csv", str(p), "--results", str(tmp_path / "kq"), "--so-lan", "50",
             "--out-json", str(tmp_path / "c.json"), "--out-md", str(tmp_path / "c.md")]
    assert m.main(chung + ["--kappa", str(tmp_path / "khong_co.json")]) == 0
    kq = json.loads((tmp_path / "c.json").read_text(encoding="utf-8"))
    assert kq["so_bo"] is True and (tmp_path / "c.md").read_text(encoding="utf-8").startswith("**SỐ SƠ BỘ")
    assert sum(kq["ket_qua"]["tat_ca"]["phobert"][g]["n"] for g in m.NHOM_CHAM) == 23
    assert sum(kq["ket_qua"]["chi_dong_y"]["phobert"][g]["n"] for g in m.NHOM_CHAM) == 22
    assert m.main(chung + ["--kappa", str(_kappa_file(tmp_path / "thap.json", 0.5))]) == 0
    assert json.loads((tmp_path / "c.json").read_text(encoding="utf-8"))["so_bo"] is True
    assert m.main(chung + ["--kappa", str(_kappa_file(tmp_path / "dat.json", 0.8))]) == 0
    assert json.loads((tmp_path / "c.json").read_text(encoding="utf-8"))["so_bo"] is False
    assert not (tmp_path / "c.md").read_text(encoding="utf-8").startswith("**SỐ SƠ BỘ")


def test_chot_bi_chan_khi_chua_co_kappa_va_ghi_dung_khi_co(tmp_path):
    p = _csv(tmp_path, moi_nhom=2)
    rows = ct.doc_csv(p)
    rows[0]["phenomenon"] = "bo"
    ct.ghi_csv(p, rows)
    test = tmp_path / "test.jsonl"
    ban_ghi = [{"uid": r["uid"], "tokens": ["a"], "aspect": r["aspect"], "label": 0} for r in rows]
    ban_ghi.append({"uid": "visfd-test-77777-GENERAL", "tokens": ["b"], "aspect": "GENERAL", "label": 2})  # không thuộc tập
    test.write_text("".join(json.dumps(b) + "\n" for b in ban_ghi), encoding="utf-8")
    out = tmp_path / "d.jsonl"
    chung = ["chot", "--csv", str(p), "--test", str(test), "--out", str(out)]
    assert m.main(chung + ["--kappa", str(tmp_path / "khong_co.json")]) == 2 and not out.exists()
    k = str(_kappa_file(tmp_path / "dat.json", 0.75))
    assert m.main(chung + ["--kappa", k]) == 0
    ra = [json.loads(d) for d in out.read_text(encoding="utf-8").splitlines()]
    assert len(ra) == 7 and all(b["phenomenon"] in m.NHOM_CHAM and b["tokens"] == ["a"] for b in ra)
    assert m.main(chung + ["--kappa", k]) == 2  # không tự ghi đè


def test_tren_ket_qua_that_o_nhom_thuong_khop_sklearn():
    goc = Path(__file__).resolve().parents[1]
    if not (goc / "results/phobert/seed42/predictions.jsonl").exists() or not (goc / ct.CSV_MAC_DINH).exists():
        pytest.skip("không có kết quả thật hoặc không có file ứng viên")
    from sklearn.metrics import accuracy_score, f1_score

    rows = [r for r in ct.doc_csv(goc / ct.CSV_MAC_DINH) if r["phenomenon"] in m.NHOM_CHAM]
    if not rows:
        pytest.skip("file ứng viên chưa có chú thích")
    _, o = m.dem_theo_nhom(rows, goc / "results", ["phobert"])
    r = m.cham_mot_mo_hinh(o["phobert"], 50, np.random.default_rng(0))
    for nhom in ("regular", "contrast"):
        uid = {x["uid"] for x in rows if x["phenomenon"] == nhom}
        acc, f1 = [], []
        for s in m.cb.SEEDS:
            u, g, d = m.cb.doc_du_doan(goc / f"results/phobert/seed{s}/predictions.jsonl")
            chon = [i for i, x in enumerate(u) if x in uid]
            acc.append(accuracy_score(g[chon], d[chon]))
            f1.append(f1_score(g[chon], d[chon], average="macro", labels=[0, 1, 2], zero_division=0))
        assert r[nhom]["n"] == len(uid)
        assert r[nhom]["accuracy"]["gia_tri"] == pytest.approx(np.mean(acc), abs=1e-12)
        assert r[nhom]["macro_f1"]["gia_tri"] == pytest.approx(np.mean(f1), abs=1e-12)


# --- Gán lại sớm, dấu hiệu, và so mô hình trong từng nhóm (GAP-032) ----------------


def test_do_trung_tach_duoc_dau_hieu_va_dem_dong_khong_dong_y():
    uid = [f"u{i}" for i in range(6)]
    h1 = ["both", "both", "negation", "contrast", "regular", "regular"]
    h2 = ["negation", "both", "negation", "contrast", "regular", "contrast"]
    k1 = ["dong_y", "negative", "positive", "dong_y", "khong_ro", "dong_y"]
    k2 = ["dong_y", "negative", "dong_y", "neutral", "dong_y", "dong_y"]
    lan1 = {u: {"phenomenon": a, "nhan_kiem_tra": b, "nhan_goc": "neutral"} for u, a, b in zip(uid, h1, k1)}
    lan2 = {u: {"phenomenon": a, "nhan_kiem_tra": b, "nhan_goc": "neutral"} for u, a, b in zip(uid, h2, k2)}
    t = m.do_trung(lan1, lan2, uid)
    assert t["so_dong"] == 6 and t["uid_lech"] == ["u0", "u5"]
    assert t["phenomenon"]["kappa"] == pytest.approx(cohen_kappa_score(h1, h2))
    assert t["phenomenon"]["giu_nguyen"]["both"] == {"lan_1": 2, "giu": 1}
    assert t["phenomenon"]["giu_nguyen"]["regular"] == {"lan_1": 2, "giu": 1}
    # `both` thành `negation` không đổi dấu hiệu phủ định, chỉ đổi dấu hiệu chuyển ý
    assert t["dau_hieu"]["phu_dinh"] == {"kappa": pytest.approx(1.0), "ti_le_trung": pytest.approx(1.0)}
    cy1, cy2 = [x in ("contrast", "both") for x in h1], [x in ("contrast", "both") for x in h2]
    assert t["dau_hieu"]["chuyen_y"]["kappa"] == pytest.approx(cohen_kappa_score(cy1, cy2))
    assert t["dau_hieu"]["chuyen_y"]["ti_le_trung"] == pytest.approx(4 / 6)
    assert t["nhan_cam_xuc"]["so_dong_lech"] == 2
    assert t["nhan_cam_xuc"]["khong_dong_y_nhan_goc"] == {"lan_1": 3, "lan_2": 2, "ca_hai": 1}


def test_kappa_gan_som_phai_khai_khoang_cach_va_ghi_ro_vao_file(tmp_path, capsys):
    p, x = _csv(tmp_path), tmp_path / "lai.xlsx"
    m.main(["rut-mau", "--csv", str(p), "--xlsx", str(x), "--moi-nhom", "5"])
    _gan_lai(x, {r["uid"]: r for r in ct.doc_csv(p)}, doi=1)
    out = tmp_path / "k.json"
    chung = ["kappa", "--csv", str(p), "--xlsx", str(x), "--out", str(out)]
    hom_nay = dt.date.today().strftime(m.DINH_DANG_NGAY)
    assert m.main(chung) == 2  # chưa đủ 14 ngày, không có cờ nào
    assert m.main(chung + ["--gan-som"]) == 1  # thiếu ngày và khoảng cách
    assert m.main(chung + ["--gan-som", "--ngay-gan-lai", hom_nay, "--khoang-cach-toi-da", "8", "--bo-qua-ngay"]) == 1
    muon = (dt.date.today() + dt.timedelta(days=m.SO_NGAY_CHO)).strftime(m.DINH_DANG_NGAY)
    assert m.main(chung + ["--gan-som", "--ngay-gan-lai", muon, "--khoang-cach-toi-da", "20"]) == 1  # không còn là gán sớm
    assert not out.exists()
    capsys.readouterr()
    assert m.main(chung + ["--gan-som", "--ngay-gan-lai", hom_nay, "--khoang-cach-toi-da", "8"]) == 0
    kq = json.loads(out.read_text(encoding="utf-8"))
    assert kq["loai"] == "nhat_quan_mot_nguoi" and kq["gan_som"] is True
    assert kq["ngay_gan_lai"] == hom_nay and kq["khoang_cach_toi_da_ngay"] == 8 and kq["so_ngay_quy_uoc"] == 14
    assert "cận trên" in capsys.readouterr().out
    assert m.kappa_da_dat(out)


def test_gop_nhom_va_so_voi_mo_hinh_moc_trong_tung_nhom(tmp_path):
    rows = ct.doc_csv(_csv(tmp_path, moi_nhom=10))
    ca_hai = {r["uid"] for r in rows if r["phenomenon"] == "both"}
    for exp in ("phobert", "asgcn"):  # asgcn sai hết nhóm cả hai, phobert đúng hết
        for s in m.cb.SEEDS:
            d = tmp_path / "kq" / exp / f"seed{s}"
            d.mkdir(parents=True)
            with (d / "predictions.jsonl").open("w", encoding="utf-8") as fh:
                for r in rows:
                    g = m.MA_NHAN[r["nhan_goc"]]
                    fh.write(json.dumps({"uid": r["uid"], "y_true": g, "y_pred": 2 if exp == "asgcn" and r["uid"] in ca_hai else g}) + "\n")
    _, o = m.dem_theo_nhom(rows, tmp_path / "kq", ["phobert", "asgcn"])
    gop = m.gop_nhom(o["asgcn"])
    assert gop.shape == (3, len(m.cb.SEEDS), 40, 9)
    r = m.cham_mot_mo_hinh(gop, 200, np.random.default_rng(0), ten_nhom=tuple(m.NHOM_GOP))
    assert {g: r[g]["n"] for g in m.NHOM_GOP} == {"regular": 10, "co_phu_dinh": 20, "co_chuyen_y": 20}
    assert r["co_phu_dinh"]["accuracy"]["gia_tri"] == pytest.approx(0.5)  # 10 dòng negation đúng, 10 dòng both sai
    s = m.so_voi_mo_hinh_moc(o, "phobert", 300, np.random.default_rng(0))
    assert list(s) == ["asgcn"]
    assert s["asgcn"]["both"]["accuracy"]["hieu"] == pytest.approx(-1.0) and s["asgcn"]["both"]["accuracy"]["p"] < 0.05
    assert s["asgcn"]["regular"]["accuracy"]["hieu"] == 0.0 and s["asgcn"]["regular"]["accuracy"]["p"] == 1.0
    x = s["asgcn"]["both"]["macro_f1"]
    assert x["hieu"] == pytest.approx(-1 / 3) and x["ci95_tren"] < 0


def test_cham_ghi_them_phan_bo_bang_gop_bang_so_moc_va_canh_bao_gan_som(tmp_path):
    p = _csv(tmp_path)
    rows = ct.doc_csv(p)
    rows[1]["nhan_kiem_tra"] = "positive"
    rows[2]["nhan_kiem_tra"] = "khong_ro"
    ct.ghi_csv(p, rows)
    _du_doan(tmp_path / "kq", rows)
    k = tmp_path / "k.json"
    k.write_text(json.dumps({"phenomenon": {"kappa": 0.733}, "so_dong": 60, "gan_som": True, "ngay_gan_lai": "13/10/2026",
                             "khoang_cach_toi_da_ngay": 8, "so_ngay_quy_uoc": 14, "uid_lech": ["x"]}), encoding="utf-8")
    assert m.main(["cham", "--csv", str(p), "--results", str(tmp_path / "kq"), "--so-lan", "50", "--kappa", str(k),
                   "--out-json", str(tmp_path / "c.json"), "--out-md", str(tmp_path / "c.md")]) == 0
    kq = json.loads((tmp_path / "c.json").read_text(encoding="utf-8"))
    md = (tmp_path / "c.md").read_text(encoding="utf-8")
    assert kq["so_bo"] is False and md.startswith("Độ nhất quán của nhãn hiện tượng: kappa 0,733 trên 60 dòng gán lại.")
    assert "không quá 8 ngày" in md and "không đủ 14 ngày" in md and "cận trên" in md
    assert kq["do_nhat_quan"]["gan_som"] is True and "uid_lech" not in kq["do_nhat_quan"]
    pb = kq["phan_bo"]["tat_ca"]
    assert sum(v["n"] for v in pb.values()) == 24 and pb["negation"]["nhan_goc"] == {"negative": 6, "neutral": 0, "positive": 0}
    assert pb["negation"]["kiem_nhan"] == {"dong_y": 4, "nhan_khac": 1, "khong_ro": 1}
    assert kq["ket_qua_gop"]["tat_ca"]["phobert"]["co_phu_dinh"]["n"] == 12
    assert set(kq["so_voi_mo_hinh_moc"]["tat_ca"]) == set(m.MO_HINH) - {"phobert"} and kq["mo_hinh_moc"] == "phobert"
    assert md.count("| Mô hình |") == 12
