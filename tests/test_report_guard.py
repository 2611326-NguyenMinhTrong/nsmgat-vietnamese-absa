"""[REQ-006] Kiểm cơ chế chống ghi đè nội dung đã sửa tay trong chuyende1/report/.

Bối cảnh: học viên lo `build_cd1_report.py` sẽ âm thầm ghi đè bản .docx họ đã
sửa tay. Các test này chứng minh: không có kịch bản nào mất nội dung mà không
có hành động tường minh (--force, kèm sao lưu tự động).
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from build_cd1_report import (  # noqa: E402
    MANIFEST_NAME,
    NGUON_SINH,
    NGUON_SUA_TAY,
    load_manifest,
    save_manifest,
    sync_all,
    sync_generated_file,
)

NAME = "test.docx"


def make_doc(text: str) -> Document:
    doc = Document()
    doc.add_paragraph(text)
    return doc


def edit_paragraph(path: Path, new_text: str) -> None:
    """Giả lập học viên gõ tay trong Word: mở file, đổi chữ, lưu lại."""
    doc = Document(str(path))
    doc.paragraphs[0].text = new_text
    doc.save(str(path))


def read_paragraph(path: Path) -> str:
    return Document(str(path)).paragraphs[0].text


# --- sync_generated_file: từng trạng thái một ---------------------------------


def test_file_chua_ton_tai_thi_ghi_moi(tmp_path):
    manifest: dict = {}
    outcome = sync_generated_file(NAME, lambda: make_doc("noi dung goc"), tmp_path, manifest)

    assert outcome == "written"
    assert (tmp_path / NAME).exists()
    assert read_paragraph(tmp_path / NAME) == "noi dung goc"
    assert manifest[NAME]["nguon"] == NGUON_SINH


def test_chua_ai_sua_thi_tu_dong_cap_nhat_theo_ban_script_moi(tmp_path):
    """Script có thể tự sửa template (vd thêm mục mới) — nếu học viên CHƯA đụng
    vào file, lần chạy sau phải áp bản mới nhất, không được coi là 'đã sửa tay'."""
    manifest: dict = {}
    sync_generated_file(NAME, lambda: make_doc("ban 1"), tmp_path, manifest)

    outcome = sync_generated_file(NAME, lambda: make_doc("ban 2 - template da doi"), tmp_path, manifest)

    assert outcome == "written"
    assert read_paragraph(tmp_path / NAME) == "ban 2 - template da doi"


def test_da_sua_tay_thi_bi_chan_va_khong_dung_vao_file(tmp_path):
    manifest: dict = {}
    sync_generated_file(NAME, lambda: make_doc("ban goc"), tmp_path, manifest)

    edit_paragraph(tmp_path / NAME, "hoc vien da sua noi dung nay")

    outcome = sync_generated_file(NAME, lambda: make_doc("ban moi cua script"), tmp_path, manifest)

    assert outcome == "blocked"
    assert read_paragraph(tmp_path / NAME) == "hoc vien da sua noi dung nay", (
        "nội dung học viên sửa phải còn nguyên sau khi bị chặn"
    )
    assert not any(tmp_path.glob("*.backup-*")), "không được tự tạo backup khi chỉ bị chặn (chưa --force)"


def test_adopt_ghi_nhan_ma_khong_dung_vao_file(tmp_path):
    manifest: dict = {}
    sync_generated_file(NAME, lambda: make_doc("ban goc"), tmp_path, manifest)
    edit_paragraph(tmp_path / NAME, "hoc vien da sua")

    outcome = sync_generated_file(NAME, lambda: make_doc("ban moi"), tmp_path, manifest, adopt=True)

    assert outcome == "adopted"
    assert read_paragraph(tmp_path / NAME) == "hoc vien da sua"
    assert manifest[NAME]["nguon"] == NGUON_SUA_TAY


def test_sau_khi_adopt_van_bi_bao_ve_vinh_vien_du_khong_doi_gi_them(tmp_path):
    """Đây là điểm cốt lõi: adopt không phải 'miễn dịch một lần' rồi lại bị ghi
    đè ở lần sau — nó phải bảo vệ VĨNH VIỄN cho tới khi có --force tường minh."""
    manifest: dict = {}
    sync_generated_file(NAME, lambda: make_doc("ban goc"), tmp_path, manifest)
    edit_paragraph(tmp_path / NAME, "hoc vien da sua")
    sync_generated_file(NAME, lambda: make_doc("khong dung"), tmp_path, manifest, adopt=True)

    # Chạy lại nhiều lần ở chế độ mặc định, không ai sửa gì thêm
    for _ in range(3):
        outcome = sync_generated_file(NAME, lambda: make_doc("ban script neu duoc ap"), tmp_path, manifest)
        assert outcome == "blocked"

    assert read_paragraph(tmp_path / NAME) == "hoc vien da sua"


def test_force_ghi_de_va_tu_dong_sao_luu_ban_da_sua(tmp_path):
    manifest: dict = {}
    sync_generated_file(NAME, lambda: make_doc("ban goc"), tmp_path, manifest)
    edit_paragraph(tmp_path / NAME, "hoc vien da sua - se bi mat")

    outcome = sync_generated_file(NAME, lambda: make_doc("ban moi cua script"), tmp_path, manifest, force=True)

    assert outcome == "forced"
    assert read_paragraph(tmp_path / NAME) == "ban moi cua script"

    backups = list(tmp_path.glob(f"{Path(NAME).stem}.backup-*{Path(NAME).suffix}"))
    assert len(backups) == 1, "phải có đúng một bản sao lưu"
    assert read_paragraph(backups[0]) == "hoc vien da sua - se bi mat", (
        "bản sao lưu phải giữ đúng nội dung đã sửa trước khi mất"
    )


def test_migrate_file_co_san_truoc_khi_co_co_che_nay(tmp_path):
    """File đã tồn tại từ trước (sinh ra trước khi cơ chế bảo vệ này được thêm
    vào), chưa có trong manifest. Nếu nội dung khớp CHÍNH XÁC bản script sẽ
    sinh ngay bây giờ -> tự nhận làm mốc, không cần hỏi, không cần ghi lại."""
    doc = make_doc("dung la ban chuan")
    doc.save(str(tmp_path / NAME))

    manifest: dict = {}
    outcome = sync_generated_file(NAME, lambda: make_doc("dung la ban chuan"), tmp_path, manifest)

    assert outcome == "baseline"
    assert manifest[NAME]["nguon"] == NGUON_SINH
    assert read_paragraph(tmp_path / NAME) == "dung la ban chuan"


def test_migrate_file_co_san_nhung_khac_ban_chuan_thi_bi_chan():
    """Ngược lại: file có sẵn không khớp bản chuẩn và chưa từng ghi nhận ->
    không đủ bằng chứng để tự tin ghi đè -> phải chặn, không suy đoán."""
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        doc = make_doc("noi dung la gi khong ro nguon goc")
        doc.save(str(tmp_path / NAME))

        manifest: dict = {}
        outcome = sync_generated_file(NAME, lambda: make_doc("ban chuan cua script"), tmp_path, manifest)

        assert outcome == "blocked"
        assert read_paragraph(tmp_path / NAME) == "noi dung la gi khong ro nguon goc"


# --- manifest persistence ------------------------------------------------------


def test_manifest_luu_va_doc_lai_dung(tmp_path):
    manifest = {NAME: {"fingerprint": "abc123", "nguon": NGUON_SINH}}
    save_manifest(tmp_path, manifest)

    assert (tmp_path / MANIFEST_NAME).exists()
    assert load_manifest(tmp_path) == manifest


def test_load_manifest_rong_khi_chua_co_file(tmp_path):
    assert load_manifest(tmp_path) == {}


# --- sync_all: hành vi tổng hợp trên cả 2 file thật của CĐ1 -------------------


def test_sync_all_lan_dau_ghi_ca_hai_file(tmp_path):
    code = sync_all(tmp_path)

    assert code == 0
    assert (tmp_path / "reference.docx").exists()
    assert (tmp_path / "ChuyenDe1_NguyenMinhTrong.docx").exists()
    assert (tmp_path / MANIFEST_NAME).exists()


def test_sync_all_chan_dung_1_file_khong_anh_huong_file_kia(tmp_path):
    """Học viên chỉ sửa cuốn báo cáo, không đụng vào reference.docx ->
    reference.docx vẫn được cập nhật bình thường, chỉ cuốn báo cáo bị chặn."""
    sync_all(tmp_path)
    edit_paragraph(tmp_path / "ChuyenDe1_NguyenMinhTrong.docx", "hoc vien dang go bai that")

    code = sync_all(tmp_path)

    assert code == 1, "phải báo lỗi vì có ít nhất 1 file bị chặn"
    assert read_paragraph(tmp_path / "ChuyenDe1_NguyenMinhTrong.docx") == "hoc vien dang go bai that"
    # reference.docx không bị sửa nên vẫn được đồng bộ bình thường, không nằm
    # trong danh sách bị chặn của manifest
    manifest = load_manifest(tmp_path)
    assert manifest["reference.docx"]["nguon"] == NGUON_SINH


def test_file_bi_khoa_khong_lam_hong_manifest(tmp_path, monkeypatch, capsys):
    """File đang mở trong Word -> PermissionError. Không được để exception thoát
    ra ngoài: nếu thoát, manifest không kịp lưu và file đã ghi xong trước đó sẽ
    bị coi là 'đã sửa tay' ở lần chạy sau -> chặn oan.

    Đây là lỗi thật gặp phải ngày 02/09/2026 khi sinh lại bìa lúc Word đang mở.
    """
    import build_cd1_report

    sync_all(tmp_path)  # lần đầu: tạo cả 2 file + manifest
    manifest_truoc = load_manifest(tmp_path)

    # Giả lập Word khoá đúng file thứ hai trong danh sách
    goc = build_cd1_report.sync_generated_file

    def khoa_file_thu_hai(name, *args, **kwargs):
        if name == "ChuyenDe1_NguyenMinhTrong.docx":
            raise PermissionError(13, "Permission denied")
        return goc(name, *args, **kwargs)

    monkeypatch.setattr(build_cd1_report, "sync_generated_file", khoa_file_thu_hai)

    code = sync_all(tmp_path)

    assert code == 1
    assert "ĐANG BỊ KHOÁ" in capsys.readouterr().out
    # Manifest vẫn phải được lưu, và mục của file ghi thành công vẫn hợp lệ
    manifest_sau = load_manifest(tmp_path)
    assert manifest_sau["reference.docx"]["nguon"] == NGUON_SINH
    assert manifest_sau["ChuyenDe1_NguyenMinhTrong.docx"] == manifest_truoc["ChuyenDe1_NguyenMinhTrong.docx"]


def test_sau_khi_mo_khoa_thi_khong_bi_chan_oan(tmp_path, monkeypatch):
    """Nối tiếp test trên: sau khi đóng Word, chạy lại phải chạy trơn, không
    được báo 'đã sửa tay' với file vừa bị khoá hụt."""
    import build_cd1_report

    sync_all(tmp_path)
    goc = build_cd1_report.sync_generated_file

    def khoa_mot_lan(name, *args, **kwargs):
        if name == "reference.docx":
            raise PermissionError(13, "Permission denied")
        return goc(name, *args, **kwargs)

    monkeypatch.setattr(build_cd1_report, "sync_generated_file", khoa_mot_lan)
    sync_all(tmp_path)

    monkeypatch.setattr(build_cd1_report, "sync_generated_file", goc)  # "đóng Word"
    code = sync_all(tmp_path)

    assert code == 0, "không được chặn oan sau khi file hết bị khoá"


def test_sync_all_bao_ve_toan_bo_khi_adopt(tmp_path):
    sync_all(tmp_path)
    edit_paragraph(tmp_path / "ChuyenDe1_NguyenMinhTrong.docx", "noi dung that cua hoc vien")

    code = sync_all(tmp_path, adopt=True)
    assert code == 0

    # Chạy lại mặc định nhiều lần sau đó — không được đụng vào nữa
    for _ in range(2):
        code = sync_all(tmp_path)
        assert code == 1

    assert read_paragraph(tmp_path / "ChuyenDe1_NguyenMinhTrong.docx") == "noi dung that cua hoc vien"
