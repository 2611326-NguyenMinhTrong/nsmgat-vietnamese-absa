"""[CD1.12k] Test scripts/kiem_sua_tay_word.py: đọc chữ file Word, so khác, và lệnh không đụng file thật."""

from __future__ import annotations

import sys
from pathlib import Path

import docx

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

import kiem_sua_tay_word as k  # noqa: E402


def _word(duong: Path, doan, hang=None) -> Path:
    d = docx.Document()
    for t in doan:
        d.add_paragraph(t)
    if hang:
        bang = d.add_table(rows=len(hang), cols=len(hang[0]))
        for i, h in enumerate(hang):
            for j, o in enumerate(h):
                bang.cell(i, j).text = o
    d.save(str(duong))
    return duong


def test_cac_khoi_bo_doan_rong_va_doc_hang_bang(tmp_path):
    p = _word(tmp_path / "a.docx", ["Một", "", "Hai"], [["x", "y"]])
    assert k.cac_khoi(p) == ["Một", "Hai", "| x | y |"]


def test_so_khac_bat_doan_sua_them_va_bo():
    nhap = ["A", "B", "C"]
    word = ["A", "B sửa", "C", "D thêm"]
    assert k.so_khac(nhap, word) == [(["B"], ["B sửa"]), ([], ["D thêm"])]
    assert k.so_khac(nhap, list(nhap)) == []


def test_chu_khac():
    assert k.chu_khac("Tính cấp thiết. Hướng", "Tính cấp thiết, hướng") == "- thiết. + thiết, - Hướng + hướng"


def test_lenh_bao_cho_sua_tay_va_khong_dung_file_that(tmp_path, monkeypatch, capsys):
    that = _word(tmp_path / "bao_cao.docx", ["Đoạn học viên đã sửa", "Đoạn chung"])
    truoc = that.read_bytes()

    def ghep_gia(argv):  # giả lệnh ghép: ghi chữ của bản nháp vào file được chỉ định
        duong = Path(argv[argv.index("--docx") + 1])
        _word(duong, ["Đoạn trong bản nháp", "Đoạn chung"])
        return 0

    monkeypatch.setattr(k.nc, "main", ghep_gia)
    ma = k.main(["--docx", str(that), "--nhap", str(tmp_path)])
    ra = capsys.readouterr().out
    assert ma == 1
    assert "WORD (sẽ mất)         : Đoạn học viên đã sửa" in ra
    assert that.read_bytes() == truoc, "file Word thật bị sửa"


def test_lenh_tra_0_khi_trung(tmp_path, monkeypatch):
    that = _word(tmp_path / "bao_cao.docx", ["Đoạn chung"])
    monkeypatch.setattr(k.nc, "main", lambda argv: (_word(Path(argv[argv.index("--docx") + 1]), ["Đoạn chung"]), 0)[1])
    assert k.main(["--docx", str(that), "--nhap", str(tmp_path)]) == 0


def _them_phuong_trinh(duong: Path, so_hieu: str) -> None:
    """Thêm một đoạn chỉ có phương trình Word (OMML) và số hiệu, như học viên tạo bằng Insert > Equation."""
    from docx.oxml import parse_xml

    d = docx.Document(str(duong))
    p = d.add_paragraph()
    p._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>x=y+1</m:t></m:r></m:oMath>'))
    p.add_run("\t" + so_hieu)
    d.save(str(duong))


def test_dong_cong_thuc_bi_bo_qua_ca_hai_phia(tmp_path):
    word = _word(tmp_path / "w.docx", ["Đoạn A"])
    _them_phuong_trinh(word, "(2.2)")
    nhap = _word(tmp_path / "n.docx", ["Đoạn A", "x = y + 1    (2.2)"])
    assert k.cac_khoi(word) == ["Đoạn A"]
    assert k.cac_khoi(nhap) == ["Đoạn A"]
    assert k.so_phuong_trinh(word) == 1 and k.so_phuong_trinh(nhap) == 0


def test_phuong_trinh_co_vai_ky_tu_lot_ra_ngoai_va_dau_cach_thua(tmp_path):
    """Thật ở file của học viên: dòng (2.2) có ")V" nằm ngoài phương trình; một đoạn có dấu cách thừa ở cuối."""
    word = _word(tmp_path / "w.docx", ["Đoạn A "])
    _them_phuong_trinh(word, ")V\t(2.2)")
    assert k.cac_khoi(word) == ["Đoạn A"]


def test_cau_van_co_so_hieu_cong_thuc_khong_bi_bo_qua(tmp_path):
    p = _word(tmp_path / "a.docx", ["Chuẩn hoá theo công thức (2.4).", "a = b (2.4)"])
    assert k.cac_khoi(p) == ["Chuẩn hoá theo công thức (2.4)."]


def test_lenh_tra_2_khi_chu_khop_nhung_co_phuong_trinh(tmp_path, monkeypatch):
    that = _word(tmp_path / "bao_cao.docx", ["Đoạn chung"])
    _them_phuong_trinh(that, "(2.2)")
    monkeypatch.setattr(k.nc, "main", lambda argv: (_word(Path(argv[argv.index("--docx") + 1]), ["Đoạn chung", "x = y + 1    (2.2)"]), 0)[1])
    assert k.main(["--docx", str(that), "--nhap", str(tmp_path)]) == 2
