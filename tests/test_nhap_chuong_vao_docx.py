"""[CD1.12] Kiểm scripts/nhap_chuong_vao_docx.py trên một cuốn Word tí hon."""

import json
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import nhap_chuong_vao_docx as m  # noqa: E402

MD3 = """# Chương 3. TÊN

Bài B trước [@B], rồi bài A [@A].
"""

MD4 = """# Chương 4. TÊN

Đoạn mở đầu trích lại [@A] và thêm [@C].

## 4.1. Mục một

Câu có số 86,64 % và [TODO: điền sau] ở giữa.

Bảng 4.1: Tên bảng

| A | B |
|---|---|
| 1 | 2 |

x = y + 1    (4.1)

### 4.1.1. Mục con

Đoạn cuối.

## Tài liệu được trích trong chương

[1] Không được vào chương.
"""

TL = {"A": "Tác giả A. Bài A. 2020.", "B": "Tác giả B. Bài B. 2021.",
      "C": "Tác giả C. Bài C. 2022. [CẦN KIỂM: chưa đối chiếu.]", "D": "Không ai trích."}


def _cuon(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Chương 3. BA", style="Heading 1")
    doc.add_paragraph("(Ngân sách: khoảng 7 trang)")
    doc.add_paragraph("")
    doc.add_paragraph("Chương 4. BỐN", style="Heading 1")
    doc.add_paragraph("(Ngân sách: khoảng 8 trang)")
    doc.add_paragraph("4.1. Cũ", style="Heading 2")
    doc.add_paragraph("[TODO: khung]")
    doc.add_paragraph("")
    doc.add_paragraph("Chương 5. NĂM", style="Heading 1")
    doc.add_paragraph("Nội dung chương năm.")
    doc.add_paragraph("TÀI LIỆU THAM KHẢO", style="Heading 1")
    doc.add_paragraph("[TODO: danh mục]")
    doc.add_paragraph("")
    doc.add_paragraph("PHỤ LỤC", style="Heading 1")
    doc.add_paragraph("Nội dung phụ lục.")
    p = tmp_path / "cuon.docx"
    doc.save(str(p))
    return p


def _nhap(tmp_path: Path, md3=MD3, md4=MD4) -> Path:
    d = tmp_path / "nhap"
    d.mkdir(exist_ok=True)
    (d / "Chuong3_X.md").write_text(md3, encoding="utf-8")
    (d / "Chuong4_Y.md").write_text(md4, encoding="utf-8")
    (d / "tai_lieu.json").write_text(json.dumps(TL, ensure_ascii=False), encoding="utf-8")
    return d


def _chay(tmp_path, p, *them, **kw):
    return m.main(["--nhap", str(_nhap(tmp_path, **kw)), "--docx", str(p), *them])


def _chu(p):
    return [(x.style.name, x.text) for x in Document(str(p)).paragraphs]


def test_doc_markdown_tach_dung_khoi_va_dung_o_muc_tai_lieu():
    loai = [k[0] for k in m.doc_markdown(MD4)]
    assert loai == ["doan", "h2", "doan", "ten", "bang", "cong_thuc", "h3", "doan"]


def test_bang_lech_so_cot_thi_nem_loi():
    with pytest.raises(ValueError):
        m.doc_markdown("| a | b |\n|---|---|\n| 1 |\n")


def test_danh_so_theo_thu_tu_xuat_hien_trong_toan_van():
    than, thu_tu = m.danh_so_trich_dan({4: MD4, 3: MD3}, TL)
    assert thu_tu == ["B", "A", "C"]  # chương 3 đi trước chương 4 dù đưa vào sau
    assert "Bài B trước [1], rồi bài A [2]." in than[3]
    assert "trích lại [2] và thêm [3]." in than[4]


def test_khoa_khong_co_trong_danh_muc_thi_nem_loi():
    with pytest.raises(KeyError):
        m.danh_so_trich_dan({1: "Câu [@KhongCo]."}, TL)


def test_hai_file_cho_mot_chuong_thi_nem_loi(tmp_path):
    d = _nhap(tmp_path)
    (d / "Chuong3_Z.md").write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        m.tim_file_chuong(d)


def test_ghep_dung_chuong_giu_nguyen_chuong_khac_va_viet_danh_muc(tmp_path):
    p = _cuon(tmp_path)
    assert _chay(tmp_path, p) == 0
    chu = _chu(p)
    ten = [t for _, t in chu]
    assert ("Normal", "Nội dung chương năm.") in chu and ("Normal", "Nội dung phụ lục.") in chu
    assert ("Heading 2", "4.1. Mục một") in chu and ("Heading 3", "4.1.1. Mục con") in chu
    assert ("Caption", "Bảng 4.1: Tên bảng") in chu
    van = "\n".join(ten)
    for cu in ("4.1. Cũ", "(Ngân sách:", "Không được vào chương", "[TODO: danh mục]", "[@", "Không ai trích"):
        assert cu not in van
    doc = Document(str(p))
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 1).text == "2"
    # thứ tự trong từng mục, và đoạn rỗng ngắt trang được giữ lại trước mục sau
    assert ten.index("Chương 3. BA") < ten.index("Bài B trước [1], rồi bài A [2].") < ten.index("Chương 4. BỐN")
    assert ten.index("Chương 4. BỐN") < ten.index("Đoạn mở đầu trích lại [2] và thêm [3].") \
        < ten.index("Đoạn cuối.") < ten.index("Chương 5. NĂM")
    assert ten[ten.index("Chương 5. NĂM") - 1] == ""
    i = ten.index("TÀI LIỆU THAM KHẢO")
    assert ten[i + 1:i + 5] == ["[1] Tác giả B. Bài B. 2021.", "[2] Tác giả A. Bài A. 2020.",
                                "[3] Tác giả C. Bài C. 2022. [CẦN KIỂM: chưa đối chiếu.]", ""]


def test_bang_nam_ngay_sau_ten_bang(tmp_path):
    p = _cuon(tmp_path)
    _chay(tmp_path, p)
    doc = Document(str(p))
    ten = next(x for x in doc.paragraphs if x.text == "Bảng 4.1: Tên bảng")
    assert ten._p.getnext() is doc.tables[0]._tbl


def test_cho_danh_dau_in_do(tmp_path):
    p = _cuon(tmp_path)
    _chay(tmp_path, p)
    para = next(x for x in Document(str(p)).paragraphs if "86,64" in x.text)
    assert [r.text for r in para.runs if r.font.color.rgb == m.DO] == ["[TODO: điền sau]"]


def test_chuong_da_co_noi_dung_thi_dung_tru_khi_ghi_de(tmp_path):
    p = _cuon(tmp_path)
    assert _chay(tmp_path, p) == 0
    truoc = p.read_bytes()
    assert _chay(tmp_path, p) == 2
    assert p.read_bytes() == truoc
    assert _chay(tmp_path, p, "--ghi-de") == 0
    doc = Document(str(p))
    assert len(doc.tables) == 1  # không nhân đôi nội dung
    assert [x.text for x in doc.paragraphs].count("[1] Tác giả B. Bài B. 2021.") == 1


def test_them_trich_dan_o_chuong_truoc_thi_chuong_sau_doi_so(tmp_path):
    p = _cuon(tmp_path)
    _chay(tmp_path, p)
    _chay(tmp_path, p, "--ghi-de", md3="# Chương 3\n\nThêm [@C] lên đầu, rồi [@B] và [@A].\n")
    ten = [t for _, t in _chu(p)]
    assert "Đoạn mở đầu trích lại [3] và thêm [1]." in ten


def test_co_sao_luu_va_chi_xem_khong_ghi(tmp_path):
    p = _cuon(tmp_path)
    truoc = p.read_bytes()
    assert _chay(tmp_path, p, "--chi-xem") == 0
    assert p.read_bytes() == truoc and not list(tmp_path.glob("*.bak"))
    _chay(tmp_path, p)
    bak = list(tmp_path.glob("cuon.*.docx.bak"))
    assert len(bak) == 1 and bak[0].read_bytes() == truoc


def _anh(tmp_path: Path) -> Path:
    from PIL import Image

    p = tmp_path / "h.png"
    Image.new("RGB", (200, 100), "white").save(p)
    return p


def test_hinh_duoc_chen_ngay_tren_ten_hinh_voi_do_rong_yeu_cau(tmp_path):
    anh = _anh(tmp_path)
    md = f"# Chương 4. TÊN\n\nĐoạn trước.\n\n![10]({anh.as_posix()})\n\nHình 4.1: Tên hình\n\n![]({anh.as_posix()})\n"
    khoi = m.doc_markdown(md)
    assert [k[0] for k in khoi] == ["doan", "hinh", "ten", "hinh"]
    assert khoi[1][1] == (anh.as_posix(), 10.0) and khoi[3][1][1] == m.RONG_HINH_CM
    cuon, nhap = _cuon(tmp_path), tmp_path / "nhap"
    nhap.mkdir()
    (nhap / "Chuong4_X.md").write_text(md, encoding="utf-8")
    (nhap / m.TEN_FILE_TAI_LIEU).write_text(json.dumps(TL), encoding="utf-8")
    assert m.main(["--nhap", str(nhap), "--docx", str(cuon), "--ghi-de"]) == 0
    doc = Document(str(cuon))
    assert len(doc.inline_shapes) == 2 and round(doc.inline_shapes[0].width.cm, 1) == 10.0
    than = m._than_muc(*m.tim_muc_cap_mot(doc, "Chương 4."))
    vi_tri = [i for i, e in enumerate(than) if e.xpath(".//w:drawing")]
    assert "".join(than[vi_tri[0] + 1].itertext()).startswith("Hình 4.1:")


def test_hinh_khong_co_tren_dia_thi_dung_khong_ghi(tmp_path, capsys):
    cuon, nhap = _cuon(tmp_path), tmp_path / "nhap"
    nhap.mkdir()
    (nhap / "Chuong4_X.md").write_text("# Chương 4. TÊN\n\n![](khong/co/hinh.png)\n\nHình 4.1: Tên\n", encoding="utf-8")
    (nhap / m.TEN_FILE_TAI_LIEU).write_text(json.dumps(TL), encoding="utf-8")
    truoc = cuon.read_bytes()
    assert m.main(["--nhap", str(nhap), "--docx", str(cuon), "--ghi-de"]) == 1
    assert "không có trên đĩa" in capsys.readouterr().out and cuon.read_bytes() == truoc


# --- Phụ lục và cờ sửa định dạng ----------------------------------------------

PL = """# PHỤ LỤC

## Phụ lục A. Bảng đầy đủ

Bảng A.1: Bảng chuyển từ thân bài

| Công trình | Năm |
|---|---|
| Bài D [@D] | 2023 |
| Bài A [@A] | 2020 |
"""


def test_phu_luc_duoc_ghep_va_trich_dan_danh_so_sau_cac_chuong(tmp_path):
    p = _cuon(tmp_path)
    nhap = _nhap(tmp_path)
    (nhap / m.TEN_FILE_PHU_LUC).write_text(PL, encoding="utf-8")
    assert m.main(["--nhap", str(nhap), "--docx", str(p)]) == 0
    doc = Document(str(p))
    chu = [x.text for x in doc.paragraphs]
    vt = chu.index("PHỤ LỤC")
    assert "Nội dung phụ lục." not in chu
    assert chu[vt + 1] == "Phụ lục A. Bảng đầy đủ" and doc.paragraphs[vt + 1].style.name == "Heading 2"
    assert chu[vt + 2] == "Bảng A.1: Bảng chuyển từ thân bài" and doc.paragraphs[vt + 2].style.name == "Caption"
    bang = doc.tables[-1]
    o = [c.text for r in bang.rows for c in r.cells]
    # D chỉ được trích ở phụ lục nên nhận số sau mọi tài liệu của các chương; A giữ số đã có từ Chương 3
    so_a = next(x for x in chu if x.endswith(TL["A"])).split("]")[0] + "]"
    so_d = next(x for x in chu if x.endswith(TL["D"])).split("]")[0] + "]"
    assert o == ["Công trình", "Năm", f"Bài D {so_d}", "2023", f"Bài A {so_a}", "2020"]
    assert int(so_d.strip("[]")) > int(so_a.strip("[]"))


def test_khong_co_file_phu_luc_thi_muc_phu_luc_giu_nguyen(tmp_path):
    p = _cuon(tmp_path)
    assert _chay(tmp_path, p) == 0
    assert "Nội dung phụ lục." in [x.text for x in Document(str(p)).paragraphs]


def test_co_sua_dinh_dang_doi_kho_giay_khong_co_thi_giu(tmp_path):
    from docx.oxml.ns import qn

    p = _cuon(tmp_path)
    _chay(tmp_path, p)
    assert int(Document(str(p)).sections[0]._sectPr.find(qn("w:pgSz")).get(qn("w:w"))) == 12240
    _chay(tmp_path, p, "--ghi-de", "--sua-dinh-dang")
    doc = Document(str(p))
    assert abs(doc.sections[0].page_width.cm - 21.0) < 0.01
    assert doc.styles["Heading 1"].element.rPr.find(qn("w:rFonts")).get(qn("w:asciiTheme")) is None


# --- Các trang trước Chương 1 (CD1.12f) ------------------------------------------

TRANG_DAU = """# TÓM TẮT

Đoạn tóm tắt có số 86,64 %.

# LỜI CAM ĐOAN

Tôi cam đoan.

>>
>> Tp. Hồ Chí Minh, ngày … tháng … năm 2026
>> NGƯỜI VIẾT
"""


def _cuon_co_trang_dau(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Tp. Hồ Chí Minh, tháng ...... năm 2026")
    doc.add_paragraph("[TODO: điền tháng]")
    for ten, than in (("LỜI CAM ĐOAN", ["Câu cũ.", "[TODO: theo mẫu]"]), ("TÓM TẮT", ["[TODO: viết sau cùng]"]),
                      ("MỤC LỤC", ["mục lục cũ", "[TODO: bấm F9]"]), ("DANH SÁCH CÁC BẢNG", ["[TODO: liệt kê]"]),
                      ("DANH SÁCH CÁC HÌNH", ["[TODO: tương tự]"]), ("DANH SÁCH CÔNG THỨC", ["[TODO: không bắt buộc]"])):
        doc.add_paragraph(ten, style="Heading 1")
        for d in than:
            doc.add_paragraph(d)
    doc.add_paragraph("")
    doc.add_paragraph("Chương 1. MỘT", style="Heading 1")
    doc.add_paragraph("(Ngân sách: khoảng 5 trang)")
    doc.add_paragraph("")
    for ten in ("Chương 3. BA", "Chương 4. BỐN"):
        doc.add_paragraph(ten, style="Heading 1")
        doc.add_paragraph("(Ngân sách: khoảng 7 trang)")
        doc.add_paragraph("")
    doc.add_paragraph("TÀI LIỆU THAM KHẢO", style="Heading 1")
    doc.add_paragraph("[TODO: danh mục]")
    p = tmp_path / "cuon_td.docx"
    doc.save(str(p))
    return p


def _chay_trang_dau(tmp_path, p, *them, trang_dau=TRANG_DAU):
    d = _nhap(tmp_path)
    (d / m.TEN_FILE_TRANG_DAU).write_text(trang_dau, encoding="utf-8")
    return m.main(["--nhap", str(d), "--docx", str(p), *them])


def _truong(el):
    from docx.oxml.ns import qn

    return ["".join(x.text or "" for x in el.iter(qn("w:instrText")))]


def test_dong_canh_giua_va_tach_muc_trang_dau():
    muc = m.doc_trang_dau(TRANG_DAU)
    assert list(muc) == ["TÓM TẮT", "LỜI CAM ĐOAN"]
    assert muc["LỜI CAM ĐOAN"] == [("doan", "Tôi cam đoan."), ("giua", ""), ("giua", "Tp. Hồ Chí Minh, ngày … tháng … năm 2026"),
                                    ("giua", "NGƯỜI VIẾT")]
    with pytest.raises(ValueError):
        m.doc_trang_dau("# A\n\nx\n\n# A\n\ny\n")
    with pytest.raises(ValueError):
        m.doc_trang_dau("chữ đứng trước tiêu đề\n\n# A\n")


def test_muc_trang_dau_da_co_noi_dung_thi_dung_tru_khi_ghi_de(tmp_path, capsys):
    p = _cuon_co_trang_dau(tmp_path)
    doc = Document(str(p))
    assert m.muc_con_la_khung(doc, "TÓM TẮT") and not m.muc_con_la_khung(doc, "LỜI CAM ĐOAN")
    assert _chay_trang_dau(tmp_path, p) == 2
    assert "LỜI CAM ĐOAN" in capsys.readouterr().out
    assert ("Normal", "Câu cũ.") in _chu(p)
    assert _chay_trang_dau(tmp_path, p, "--ghi-de") == 0


def test_ghep_trang_dau_thay_muc_dung_danh_sach_va_bo_cho_danh_dau(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = _cuon_co_trang_dau(tmp_path)
    assert _chay_trang_dau(tmp_path, p, "--ghi-de") == 0
    doc = Document(str(p))
    chu = [(x.style.name, x.text) for x in doc.paragraphs]
    dau = chu[:chu.index(("Heading 1", "Chương 1. MỘT"))]
    # nội dung của hai mục có trong TrangDau.md đã vào, câu cũ và mọi dòng [TODO ...] ở trang đầu đã mất
    assert ("Normal", "Đoạn tóm tắt có số 86,64 %.") in dau and ("Normal", "Tôi cam đoan.") in dau
    assert ("Normal", "Câu cũ.") not in dau and not [t for _, t in dau if "[TODO" in t]
    assert ("Normal", "Tp. Hồ Chí Minh, tháng ...... năm 2026") in dau  # trang bìa chỉ mất dòng đánh dấu
    assert ("Normal", "mục lục cũ") in dau  # mục không có trong TrangDau.md thì giữ, chỉ bỏ dòng đánh dấu
    # mục ngoài quy định bị bỏ; các mục còn lại theo đúng thứ tự cũ
    assert [t for s, t in dau if s == "Heading 1"] == ["LỜI CAM ĐOAN", "TÓM TẮT", "MỤC LỤC", "DANH SÁCH CÁC BẢNG", "DANH SÁCH CÁC HÌNH"]
    # khối chữ ký canh giữa
    ky = next(x for x in doc.paragraphs if x.text == "NGƯỜI VIẾT")
    assert ky.alignment == WD_ALIGN_PARAGRAPH.CENTER
    # mỗi mục trước Chương 1 ở một trang mới, trừ mục đầu
    h1 = [x for x in doc.paragraphs if x.style.name == "Heading 1"]
    truoc = h1[:[x.text for x in h1].index("Chương 1. MỘT")]
    assert [bool(x.paragraph_format.page_break_before) for x in truoc] == [False, True, True, True, True]
    assert not h1[len(truoc)].paragraph_format.page_break_before
    # hai danh sách: một dòng đầu cột rồi một trường TOC đọc các trường TC theo mã
    for ten, dau_cot, ma in (("DANH SÁCH CÁC BẢNG", "BẢNG", "B"), ("DANH SÁCH CÁC HÌNH", "HÌNH", "H")):
        than = m._than_muc(*m.tim_muc_cap_mot(doc, ten))
        chu_than = ["".join(t.text or "" for t in e.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")) for e in than]
        assert chu_than[0] == f"{dau_cot}TRANG"  # dấu tab nằm ở nút w:tab, không ở w:t
        assert _truong(than[1]) == [rf"TOC \f {ma} \h \z"]


def test_ten_bang_mang_truong_tc_khong_an(tmp_path):
    p = _cuon_co_trang_dau(tmp_path)
    assert _chay_trang_dau(tmp_path, p, "--ghi-de") == 0
    doc = Document(str(p))
    ten = next(x for x in doc.paragraphs if x.text == "Bảng 4.1: Tên bảng")  # chữ hiện ra không đổi
    assert _truong(ten._p) == [r' TC "Bảng 4.1: Tên bảng" \f B \l 1 ']
    # trường TC mang w:vanish thì Word không gom vào danh sách (đã thử trên Word ngày 13/10/2026)
    assert not ten._p.xpath(".//w:vanish")


def test_khong_co_file_trang_dau_thi_trang_dau_giu_nguyen(tmp_path):
    p = _cuon_co_trang_dau(tmp_path)
    assert m.main(["--nhap", str(_nhap(tmp_path)), "--docx", str(p)]) == 0
    chu = _chu(p)
    assert ("Normal", "[TODO: viết sau cùng]") in chu and ("Heading 1", "DANH SÁCH CÔNG THỨC") in chu
    assert not [x for x in Document(str(p)).paragraphs if x.paragraph_format.page_break_before]


def test_bo_muc_giu_doan_mang_ngat_section(tmp_path):
    from docx.enum.section import WD_SECTION

    doc = Document()
    doc.add_paragraph("A", style="Heading 1")
    doc.add_paragraph("thân A")
    doc.add_paragraph("B", style="Heading 1")
    doc.add_paragraph("thân B")
    doc.add_section(WD_SECTION.NEW_PAGE)  # ngắt section nằm trong đoạn cuối của mục B
    doc.add_paragraph("Chương 1. MỘT", style="Heading 1")
    assert m.bo_muc(doc, "B") and not m.bo_muc(doc, "KHÔNG CÓ")
    assert [x.text for x in doc.paragraphs if x.style.name == "Heading 1"] == ["A", "Chương 1. MỘT"]
    assert len(doc.sections) == 2 and "thân B" not in [x.text for x in doc.paragraphs]


def test_trang_bia_duoc_thu_gon_dong_trong_dong_co_chu_giu_nguyen():
    doc = Document()
    doc.add_paragraph("TÊN TRƯỜNG")
    for _ in range(5):
        doc.add_paragraph("")
    doc.add_paragraph("CHUYÊN ĐỀ")
    for _ in range(2):
        doc.add_paragraph("")
    doc.add_paragraph("Tp. Hồ Chí Minh")
    doc.add_paragraph("LỜI CẢM ƠN", style="Heading 1")
    for _ in range(6):
        doc.add_paragraph("")  # dòng trống ngoài trang bìa không bị đụng tới
    assert m.thu_gon_trang_bia(doc) == 5 - m.SO_DONG_TRONG_BIA
    chu = [x.text for x in doc.paragraphs]
    assert chu[:chu.index("LỜI CẢM ƠN")] == ["TÊN TRƯỜNG", "", "", "", "CHUYÊN ĐỀ", "", "", "Tp. Hồ Chí Minh"]
    assert len(chu) - chu.index("LỜI CẢM ƠN") - 1 == 6
    trong_bia = [x for x in doc.paragraphs[:chu.index("LỜI CẢM ƠN")] if not x.text]
    assert all(x.paragraph_format.line_spacing == 1.0 and x.paragraph_format.space_after == 0 for x in trong_bia)
    assert m.thu_gon_trang_bia(doc) == 0  # chạy lại không bỏ thêm gì
