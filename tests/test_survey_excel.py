"""[09/09/2026] Test cầu nối Excel của survey_tools.py.

Vì sao cần: học viên phải sửa `survey_matrix.csv` bằng tay (điền 3 cột về số
công bố của từng bài báo), mà mở thẳng file .csv bằng Excel trên máy Windows
tiếng Việt cho ra hai lỗi cùng lúc:

  1. Chữ tiếng Việt thành rác: "[CẦN TÌM]" -> "[Cáº¦N TÃŒM]"
     (file UTF-8 không BOM, Excel đoán nhầm là mã ANSI của hệ thống)
  2. Cả dòng nằm gọn trong cột A
     (dấu tách danh sách của Windows locale tiếng Việt là ";" chứ không phải ",")

Nên có vòng xuất/nhập qua .xlsx. Điều PHẢI đúng: vòng đó **không được làm sai
lệch nội dung** — nếu nó âm thầm nuốt một ô hay đổi một ký tự, dữ liệu khảo sát
của cả 15 tuần sẽ hỏng mà không ai biết.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

pytest.importorskip("openpyxl")

import survey_tools  # noqa: E402

CSV_THAT = REPO_ROOT / "chuyende1" / "survey" / "survey_matrix.csv"


def doc_tho(path: Path) -> list[list[str]]:
    """Đọc thô, không qua load_rows — để so sánh từng ô đúng như trên đĩa."""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return [r for r in csv.reader(fh) if any(c.strip() for c in r)]


@pytest.fixture
def ban_sao(tmp_path) -> Path:
    """Bản sao của file khảo sát THẬT — test trên dữ liệu thật, không phải
    dữ liệu bịa, vì chính dữ liệu thật mới có dấu tiếng Việt, dấu phẩy trong
    ô ghi_chu, và ô rỗng."""
    p = tmp_path / "survey_matrix.csv"
    p.write_bytes(CSV_THAT.read_bytes())
    return p


def test_vong_xuat_nhap_GIU_NGUYEN_tung_o(ban_sao):
    """Xuất ra .xlsx rồi nhập ngược lại phải cho ra đúng từng ô như cũ."""
    truoc = doc_tho(ban_sao)

    assert survey_tools.main(["excel", "--csv", str(ban_sao)]) == 0
    assert ban_sao.with_suffix(".xlsx").exists()
    assert survey_tools.main(["tu-excel", "--csv", str(ban_sao)]) == 0

    sau = doc_tho(ban_sao)
    assert len(sau) == len(truoc), "số dòng đổi sau khi đi một vòng"
    for i, (a, b) in enumerate(zip(truoc, sau)):
        assert a == b, f"dòng {i + 1} bị đổi:\n  trước: {a}\n  sau  : {b}"


def test_sau_khi_nhap_van_doc_duoc_bang_load_rows(ban_sao):
    """BOM phải bị bỏ qua, nếu không tên cột đầu thành '\\ufeffref_key' và mọi
    công cụ khác gãy."""
    survey_tools.main(["excel", "--csv", str(ban_sao)])
    survey_tools.main(["tu-excel", "--csv", str(ban_sao)])

    rows = survey_tools.load_rows(ban_sao)
    assert rows, "không đọc được dòng nào"
    assert "ref_key" in rows[0], f"tên cột bị hỏng: {list(rows[0])[:3]}"
    assert rows[0]["ref_key"], "ô đầu tiên rỗng"


def test_file_csv_ghi_ra_co_BOM_de_excel_doc_dung_tieng_viet(ban_sao):
    """Không có BOM thì Excel đọc UTF-8 như ANSI -> '[CẦN TÌM]' thành rác."""
    survey_tools.main(["excel", "--csv", str(ban_sao)])
    survey_tools.main(["tu-excel", "--csv", str(ban_sao)])
    assert ban_sao.read_bytes()[:3] == b"\xef\xbb\xbf"


def test_TU_CHOI_nhap_khi_header_bi_doi(ban_sao):
    """Học viên lỡ tay đổi tên cột hoặc chèn cột trong Excel thì phải dừng
    lại, KHÔNG được ghi đè file .csv bằng dữ liệu lệch cột."""
    from openpyxl import load_workbook

    survey_tools.main(["excel", "--csv", str(ban_sao)])
    xlsx = ban_sao.with_suffix(".xlsx")
    wb = load_workbook(xlsx)
    wb.active.cell(row=1, column=2).value = "ten_cot_bi_doi"
    wb.save(xlsx)

    truoc = ban_sao.read_bytes()
    assert survey_tools.main(["tu-excel", "--csv", str(ban_sao)]) == 1
    assert ban_sao.read_bytes() == truoc, "đã ghi đè .csv dù header sai"


def test_o_de_trong_van_giu_nguyen_la_rong(ban_sao, tmp_path):
    """openpyxl trả None cho ô rỗng. Nếu không xử lý, ô rỗng sẽ thành chuỗi
    'None' — hỏng dữ liệu một cách rất khó phát hiện."""
    survey_tools.main(["excel", "--csv", str(ban_sao)])
    survey_tools.main(["tu-excel", "--csv", str(ban_sao)])

    noi_dung = ban_sao.read_text(encoding="utf-8-sig")
    assert ",None," not in noi_dung and not noi_dung.rstrip().endswith(",None")
