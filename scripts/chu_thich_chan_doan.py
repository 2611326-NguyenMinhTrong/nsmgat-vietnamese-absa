"""[CD1.8b] Chú thích tập chẩn đoán bằng Excel, không mở thẳng file CSV.

Vì sao cần: data/diagnostic/candidates.csv dùng dấu phẩy, máy Windows tiếng Việt
mặc định dấu chấm phẩy nên bấm đúp vào CSV là dính cột; Excel còn tự đổi kiểu dữ
liệu khi lưu lại. File này là 20 đến 25 giờ làm tay, nên đi đường vòng qua .xlsx:

    python scripts/chu_thich_chan_doan.py excel       # CSV -> .xlsx có ô chọn sẵn giá trị
    python scripts/chu_thich_chan_doan.py tu-excel    # .xlsx -> CSV, kiểm trước khi ghi
    python scripts/chu_thich_chan_doan.py tien-do     # đã gán bao nhiêu, còn thiếu nhóm nào

Hướng dẫn gán nhãn: docs/annotation_guideline.md. Hai bộ giá trị hợp lệ dưới đây
là nguồn duy nhất; hướng dẫn chép theo file này.

Thứ tự dòng trong .xlsx KHÔNG phải thứ tự trong CSV: các câu được xáo trong từng
nhóm ứng viên (theo --seed) rồi xếp xen kẽ bốn nhóm, mọi khía cạnh của một câu nằm
liền nhau. Nhờ vậy dừng ở dòng nào thì phần đã gán cũng trải đều bốn nhóm. CSV vẫn
giữ nguyên thứ tự cũ; hai file ghép với nhau theo uid.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import random
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Dict, List, Sequence

sys.path.insert(0, str(Path(__file__).parent))

from build_diagnostic_candidates import COT, COT_CHU_THICH, NHOM, TEN_LOP, ma_cau  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CSV_MAC_DINH = Path("data/diagnostic/candidates.csv")
XLSX_MAC_DINH = Path("data/diagnostic/candidates_chu_thich.xlsx")
TEN_SHEET = "chu_thich"

# Cột `phenomenon`. `bo` = dòng không dùng được (câu vô nghĩa, không tìm thấy ý kiến
# nào về khía cạnh). Định nghĩa từng giá trị: docs/annotation_guideline.md mục 3.
HIEN_TUONG = ("negation", "contrast", "both", "regular", "bo")
# Cột `nhan_kiem_tra`. `dong_y` = nhãn gốc đúng; một tên lớp = nhãn mà người chú thích
# cho là đúng thay cho nhãn gốc; `khong_ro` = không quyết được.
NHAN_KIEM_TRA = ("dong_y",) + TEN_LOP + ("khong_ro",)
HOP_LE = {"phenomenon": HIEN_TUONG, "nhan_kiem_tra": NHAN_KIEM_TRA}
CHI_TIEU_MOI_NHOM = 100
NHOM_CAN_DU = ("negation", "contrast", "regular")

COT_XLSX = ("stt", "uid", "aspect", "text", "nhan_goc", "tu_khoa_kich_hoat") + COT_CHU_THICH
RONG_COT = {"stt": 6, "uid": 26, "aspect": 14, "text": 90, "nhan_goc": 11, "tu_khoa_kich_hoat": 18,
            "phenomenon": 13, "nhan_kiem_tra": 14, "ghi_chu": 40}

Dong = Dict[str, str]


def doc_csv(duong_dan: Path) -> List[Dong]:
    with duong_dan.open(encoding="utf-8-sig", newline="") as fh:
        doc = csv.DictReader(fh)
        if tuple(doc.fieldnames or ()) != COT:
            raise ValueError(f"{duong_dan} không đúng mười cột chuẩn. Có: {doc.fieldnames}")
        return [dict(r) for r in doc]


def ghi_csv(duong_dan: Path, rows: Sequence[Dong]) -> None:
    with duong_dan.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COT, lineterminator="\r\n")
        w.writeheader()
        w.writerows(rows)


def thu_tu_chu_thich(rows: Sequence[Dong], seed: int) -> List[Dong]:
    """Xáo câu trong từng nhóm ứng viên, xếp xen kẽ bốn nhóm, giữ các khía cạnh của một câu liền nhau."""
    theo_cau: Dict[str, List[Dong]] = {}
    for r in rows:
        theo_cau.setdefault(ma_cau(r["uid"]), []).append(r)
    hang_doi: Dict[str, List[str]] = {g: [] for g in NHOM}
    for cau, ds in sorted(theo_cau.items()):
        hang_doi[ds[0]["nhom_ung_vien"]].append(cau)
    rng = random.Random(seed)
    for g in NHOM:
        rng.shuffle(hang_doi[g])
    ra: List[Dong] = []
    while any(hang_doi.values()):
        for g in NHOM:
            if hang_doi[g]:
                ra.extend(sorted(theo_cau[hang_doi[g].pop()], key=lambda r: r["uid"]))
    return ra


def loi_gia_tri(rows: Sequence[Dong]) -> List[str]:
    """Các ô chú thích có giá trị ngoài bộ hợp lệ, hoặc dòng mới điền một trong hai ô bắt buộc."""
    loi: List[str] = []
    for r in rows:
        for cot, hop_le in HOP_LE.items():
            v = (r.get(cot) or "").strip()
            if v and v not in hop_le:
                loi.append(f"{r['uid']}: {cot} = '{v}' không hợp lệ (chỉ nhận {', '.join(hop_le)})")
        co = [bool((r.get(c) or "").strip()) for c in HOP_LE]
        if any(co) and not all(co):
            loi.append(f"{r['uid']}: mới điền một trong hai ô phenomenon và nhan_kiem_tra")
    return loi


def da_xong(r: Dong) -> bool:
    return all((r.get(c) or "").strip() for c in HOP_LE)


# --- Ba lệnh con -----------------------------------------------------------------


def cmd_excel(csv_path: Path, xlsx: Path, seed: int, ghi_de: bool) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    if xlsx.exists() and not ghi_de:
        print(f"DỪNG: {xlsx} đã tồn tại. Nếu trong đó có chú thích chưa đưa về CSV thì chạy `tu-excel` trước.\n"
              f"Chắc chắn muốn tạo lại từ CSV: thêm --ghi-de (file .xlsx cũ được sao lưu).")
        return 2
    rows = doc_csv(csv_path)
    loi = loi_gia_tri(rows)
    if loi:
        print("CSV đang có giá trị không hợp lệ, sửa trước khi xuất:\n  " + "\n  ".join(loi[:20]))
        return 1
    if xlsx.exists():
        shutil.copy2(xlsx, xlsx.with_name(f"{xlsx.stem}.{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx.bak"))

    wb = Workbook()
    ws = wb.active
    ws.title = TEN_SHEET
    ws.append(list(COT_XLSX))
    for i, r in enumerate(thu_tu_chu_thich(rows, seed), start=1):
        ws.append([i] + [r[c] for c in COT_XLSX[1:]])
    xam = PatternFill("solid", fgColor="EDEDED")
    vang = PatternFill("solid", fgColor="FFF2CC")
    for o in ws[1]:
        o.font = Font(bold=True)
        o.fill = xam
    for hang in ws.iter_rows(min_row=2):
        for o, cot in zip(hang, COT_XLSX):
            o.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(o.value, str):
                o.data_type = "s"  # câu bắt đầu bằng "=" hay "+" không được thành công thức
            duoc_sua = cot in COT_CHU_THICH
            o.protection = Protection(locked=not duoc_sua)
            if duoc_sua:
                o.fill = vang
    for i, cot in enumerate(COT_XLSX, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = RONG_COT[cot]
    ws.freeze_panes = "E2"
    cuoi = len(rows) + 1
    for cot, hop_le in HOP_LE.items():
        chu = ws.cell(row=1, column=COT_XLSX.index(cot) + 1).column_letter
        dv = DataValidation(type="list", formula1=f'"{",".join(hop_le)}"', allow_blank=True,
                            errorTitle="Giá trị không hợp lệ", error=f"Chỉ được một trong: {', '.join(hop_le)}")
        ws.add_data_validation(dv)
        dv.add(f"{chu}2:{chu}{cuoi}")
    # Khoá các cột dữ liệu để không sửa nhầm câu hay nhãn gốc. Không đặt mật khẩu:
    # mục đích là chống lỡ tay, không phải chống người dùng.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    print(f"Đã ghi {xlsx}: {len(rows)} dòng, {sum(da_xong(r) for r in rows)} dòng đã có chú thích.\n"
          f"Chỉ ba cột màu vàng sửa được. Gán xong một đợt: lưu, đóng Excel, chạy `tu-excel`.")
    return 0


def cmd_tu_excel(xlsx: Path, csv_path: Path) -> int:
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True, data_only=True)[TEN_SHEET]
    hang = ws.iter_rows(values_only=True)
    tieu_de = tuple(str(x) for x in next(hang))
    if tieu_de != COT_XLSX:
        print(f"DỪNG: dòng tiêu đề của {xlsx} đã bị đổi. Có: {tieu_de}")
        return 1
    moi: Dict[str, Dong] = {}
    for h in hang:
        r = {c: ("" if v is None else str(v).strip()) for c, v in zip(COT_XLSX, h)}
        if r["uid"]:
            if r["uid"] in moi:
                print(f"DỪNG: uid {r['uid']} xuất hiện hai lần trong {xlsx}")
                return 1
            moi[r["uid"]] = r
    # Ghi lại đúng nhãn gốc vào ô kiểm tra nghĩa là "đồng ý": quy về một cách viết để
    # tỉ lệ bất đồng với nhãn gốc không bị đếm sai.
    quy_ve = [r["uid"] for r in moi.values() if r["nhan_kiem_tra"] and r["nhan_kiem_tra"] == r["nhan_goc"]]
    for uid in quy_ve:
        moi[uid]["nhan_kiem_tra"] = "dong_y"
    if quy_ve:
        print(f"{len(quy_ve)} dòng ghi nhan_kiem_tra trùng nhãn gốc, đã quy về dong_y khi ghi CSV "
              f"(file .xlsx không bị sửa): {', '.join(u.replace('visfd-test-', '') for u in quy_ve[:10])}"
              + (" …" if len(quy_ve) > 10 else ""))
    rows = doc_csv(csv_path)
    loi: List[str] = []
    if set(moi) != {r["uid"] for r in rows}:
        loi.append(f"tập uid của .xlsx ({len(moi)}) không trùng với CSV ({len(rows)})")
    else:
        for r in rows:
            m = moi[r["uid"]]
            for c in ("aspect", "text", "nhan_goc"):
                if m[c] != r[c].strip():
                    loi.append(f"{r['uid']}: cột {c} trong .xlsx khác CSV (cột này không được sửa)")
        loi += loi_gia_tri(list(moi.values()))
    if loi:
        print(f"DỪNG, chưa ghi gì. {len(loi)} lỗi:\n  " + "\n  ".join(loi[:30]))
        return 1
    doi = 0
    for r in rows:
        for c in COT_CHU_THICH:
            if (r[c] or "").strip() != moi[r["uid"]][c]:
                r[c] = moi[r["uid"]][c]
                doi += 1
    if not doi:
        print("Không có ô nào đổi so với CSV. Không ghi gì.")
        return 0
    ban = csv_path.with_name(f"{csv_path.stem}.{dt.datetime.now():%Y%m%d-%H%M%S}.csv.bak")
    shutil.copy2(csv_path, ban)
    ghi_csv(csv_path, rows)
    print(f"Đã cập nhật {doi} ô trong {csv_path}. Bản sao lưu: {ban}")
    return cmd_tien_do(csv_path)


def cmd_tien_do(csv_path: Path) -> int:
    rows = doc_csv(csv_path)
    loi = loi_gia_tri(rows)
    xong = [r for r in rows if da_xong(r)]
    print(f"Đã gán {len(xong)} / {len(rows)} dòng.")
    bang = Counter((r["phenomenon"], r["nhan_goc"]) for r in xong)
    print(f"{'hiện tượng':<10} " + " ".join(f"{t:>9}" for t in TEN_LOP) + f" {'tổng':>6} {'còn thiếu':>10}")
    for ht in HIEN_TUONG:
        so = [bang[(ht, t)] for t in TEN_LOP]
        thieu = max(0, CHI_TIEU_MOI_NHOM - sum(so)) if ht in NHOM_CAN_DU else ""
        print(f"{ht:<10} " + " ".join(f"{s:>9}" for s in so) + f" {sum(so):>6} {thieu:>10}")
    kt = Counter(r["nhan_kiem_tra"] for r in xong)
    print("Kiểm nhãn gốc: " + ", ".join(f"{k}={kt[k]}" for k in NHAN_KIEM_TRA))
    if loi:
        print(f"{len(loi)} lỗi giá trị:\n  " + "\n  ".join(loi[:20]))
        return 1
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Chú thích tập chẩn đoán qua Excel.")
    sub = ap.add_subparsers(dest="lenh", required=True)
    p = sub.add_parser("excel", help="CSV -> .xlsx để chú thích")
    p.add_argument("--csv", type=Path, default=CSV_MAC_DINH)
    p.add_argument("--xlsx", type=Path, default=XLSX_MAC_DINH)
    p.add_argument("--seed", type=int, default=42, help="hạt giống cho thứ tự dòng trong .xlsx (mặc định 42)")
    p.add_argument("--ghi-de", action="store_true", help="tạo lại .xlsx dù đã tồn tại (sao lưu file cũ)")
    p = sub.add_parser("tu-excel", help=".xlsx -> CSV, kiểm trước khi ghi")
    p.add_argument("--xlsx", type=Path, default=XLSX_MAC_DINH)
    p.add_argument("--csv", type=Path, default=CSV_MAC_DINH)
    p = sub.add_parser("tien-do", help="đếm số dòng đã gán theo hiện tượng và nhãn")
    p.add_argument("--csv", type=Path, default=CSV_MAC_DINH)
    a = ap.parse_args(argv)
    if a.lenh == "excel":
        return cmd_excel(a.csv, a.xlsx, a.seed, a.ghi_de)
    if a.lenh == "tu-excel":
        return cmd_tu_excel(a.xlsx, a.csv)
    return cmd_tien_do(a.csv)


if __name__ == "__main__":
    raise SystemExit(main())
