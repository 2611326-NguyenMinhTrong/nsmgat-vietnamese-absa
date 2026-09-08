"""Chuyển Markdown sang Word (.docx) theo quy định trình bày của UTE.

[REQ-004] Định dạng lấy từ scripts/ute_docx.py — xem module đó để biết nguồn quy định.

Cú pháp Markdown hỗ trợ (cố ý giới hạn, đủ cho tài liệu học vụ ngắn):
  # ## ### ####    tiêu đề
  **đậm**  *nghiêng*  `mã`
  - mục          gạch đầu dòng
  - [ ] mục      ô đánh dấu (in ra ☐ để tích tay trên bản giấy)
  1. mục         danh sách đánh số
  > trích dẫn
  | a | b |      bảng (cần dòng phân cách ---)
  ---            đường kẻ ngang

KHÔNG dùng cho cuốn chuyên đề 30 trang: tài liệu đó cần mục lục tự động, chú thích
bảng/hình đánh số theo chương và tham chiếu chéo — Word làm tốt hơn nhiều so với
việc sinh từ Markdown. Xem scripts/build_cd1_report.py.

Dùng:
    python scripts/md_to_docx_ute.py <input.md> <output.docx>
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))

from ute_docx import (  # noqa: E402
    FONT,
    HEADING_SIZES,
    LINE_SPACING,
    SIZE,
    TABLE_TEXT_SIZE,
    new_document,
    set_page_numbering,
)

# Console Windows mac dinh la cp1252, khong in duoc chu tieng Viet co dau ->
# UnicodeEncodeError. Ep stdout ve UTF-8 ngay dau chuong trinh.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# --- Định dạng trong dòng: **đậm**, *nghiêng*, `mã` --------------------------

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def add_inline(paragraph, text: str) -> None:
    for piece in INLINE_PATTERN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*"):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        if run.font.name != "Consolas":
            run.font.name = FONT
        run.font.size = SIZE


# --- Khối ---------------------------------------------------------------------


def add_heading(doc, level: int, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = LINE_SPACING
    run = para.add_run(re.sub(r"[*`]", "", text))
    run.bold = True
    run.font.name = FONT
    run.font.size = HEADING_SIZES.get(level, SIZE)


def add_list_item(doc, text: str, numbered: bool) -> None:
    # "- [ ]" -> ô vuông, để GVHD tích tay trên bản in
    checkbox = re.match(r"\[([ xX])\]\s*(.*)", text)
    if checkbox:
        text = f"{'☒' if checkbox.group(1) in 'xX' else '☐'} {checkbox.group(2)}"
        style = None
    else:
        style = "List Number" if numbered else "List Bullet"

    para = doc.add_paragraph(style=style)
    if style is None:
        para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.line_spacing = LINE_SPACING
    para.paragraph_format.space_after = Pt(3)
    add_inline(para, text)


def add_quote(doc, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(1.0)
    para.paragraph_format.line_spacing = LINE_SPACING
    add_inline(para, text)
    for run in para.runs:
        run.italic = True


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"

    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row[: len(rows[0])]):
            para = table.cell(r, c).paragraphs[0]
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, cell_text)
            for run in para.runs:
                run.font.size = TABLE_TEXT_SIZE
                if r == 0:
                    run.bold = True
    doc.add_paragraph()


def add_rule(doc) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(key), value)
    border.append(bottom)
    para._p.get_or_add_pPr().append(border)


# --- Bộ chuyển ----------------------------------------------------------------


def convert(md_path: Path, docx_path: Path) -> None:
    doc = new_document()
    # new_document() để section đầu KHÔNG đánh số trang, vì nó được thiết kế cho
    # tài liệu có trang bìa. Tài liệu sinh từ Markdown không có bìa nên phải bật lại.
    set_page_numbering(doc.sections[0], fmt="decimal", start=1)

    lines = md_path.read_text(encoding="utf-8").splitlines()

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        # Bảng: "| ... |" theo sau bởi dòng phân cách "|---|"
        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and re.fullmatch(r"\|[\s:|-]+\|", lines[i + 1].strip() or "")
        ):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            add_table(doc, rows)
            continue

        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            add_rule(doc)
        elif heading := re.match(r"(#{1,6})\s+(.*)", stripped):
            add_heading(doc, len(heading.group(1)), heading.group(2))
        elif stripped.startswith(">"):
            add_quote(doc, stripped.lstrip("> ").strip())
        elif bullet := re.match(r"[-*+]\s+(.*)", stripped):
            add_list_item(doc, bullet.group(1), numbered=False)
        elif numbered := re.match(r"\d+\.\s+(.*)", stripped):
            add_list_item(doc, numbered.group(1), numbered=True)
        else:
            para = doc.add_paragraph()
            para.paragraph_format.line_spacing = LINE_SPACING
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(para, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"OK -> {sys.argv[2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
