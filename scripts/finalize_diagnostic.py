"""[CD1.8d / S6.4] Đo độ nhất quán của nhãn tập chẩn đoán bằng cách gán lại một mẫu.

Chưa có người chú thích thứ hai, nên dùng phương án đã ghi trong plan: chính người chú
thích gán lại một mẫu sau ít nhất hai tuần, không xem nhãn cũ. Con số thu được là độ
NHẤT QUÁN CỦA MỘT NGƯỜI (intra-annotator), không được gọi là độ đồng thuận giữa hai người.
Nếu sau này có người thứ hai thì đưa đúng file .xlsx này cho họ gán; lệnh `kappa` không đổi.

    python scripts/finalize_diagnostic.py rut-mau     # rút 60 dòng, ghi file .xlsx TRỐNG nhãn để gán lại
    python scripts/finalize_diagnostic.py kappa       # sau khi gán lại xong: tính kappa, in các dòng lệch
    python scripts/finalize_diagnostic.py cham        # chấm các mô hình theo nhóm hiện tượng
    python scripts/finalize_diagnostic.py chot        # ghi file .jsonl của tập, chỉ khi kappa đã đạt

Mẫu được rút PHÂN TẦNG theo nhãn hiện tượng của lần gán đầu, mỗi nhóm số dòng bằng nhau
(mặc định 15 x 4 nhóm), không rút theo tỉ lệ thật. Lý do: rút theo tỉ lệ thật thì nhóm
`both` chỉ được khoảng 6 dòng, không đủ để biết quy tắc của nhóm đó có lặp lại được không.
Khi báo cáo phải nói rõ mẫu là phân tầng.

`cham` KHÔNG chạy mô hình nào. Mọi dòng của tập chẩn đoán là một mẫu của tập kiểm tra, nên dự
đoán của nó đã nằm sẵn trong results/<exp>/seed<N>/predictions.jsonl; lệnh chỉ ghép dự đoán
đó với nhãn hiện tượng. Khoảng tin cậy lấy bằng bootstrap theo CÂU như compare_baselines.py,
vì một câu có thể góp dòng cho nhiều nhóm. Chừng nào chưa có data/diagnostic/kappa.json đạt
ngưỡng thì kết quả được đánh dấu SƠ BỘ và không được đưa vào báo cáo.

Nhãn dùng để chấm là NHÃN GỐC của bộ dữ liệu, để so được với bảng kết quả chính. Lệnh chấm
thêm một lần trên riêng các dòng người chú thích đồng ý với nhãn gốc, để biết kết luận có
phụ thuộc vào các dòng nhãn gốc đáng ngờ hay không.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import random
import sys
from collections import Counter
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).parent))

import chu_thich_chan_doan as ct  # noqa: E402
import compare_baselines as cb  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

XLSX_MAC_DINH = Path("data/diagnostic/gan_lai_60.xlsx")
SO_NGAY_CHO = 14
NGUONG_KAPPA = 0.70
NHOM_RUT = ("negation", "contrast", "both", "regular")
DINH_DANG_NGAY = "%d/%m/%Y"
KAPPA_MAC_DINH = Path("data/diagnostic/kappa.json")
MO_HINH = ("lexicon", "bilstm", "phobert", "asgcn", "asgcn_linked", "senticgcn")
NHOM_CHAM = ("regular", "negation", "contrast", "both")  # nhóm đầu là mốc so
TEN_NHOM = {"regular": "Thường", "negation": "Phủ định", "contrast": "Chuyển ý", "both": "Cả hai"}
MA_NHAN = {"negative": 0, "neutral": 1, "positive": 2}

Dong = Dict[str, str]


def file_so(xlsx: Path) -> Path:
    return xlsx.with_suffix(".json")


def rut_phan_tang(rows: Sequence[Dong], moi_nhom: int, seed: int) -> List[Dong]:
    """Rút `moi_nhom` dòng cho mỗi nhãn hiện tượng, rồi xáo thứ tự chung. Tất định theo seed."""
    rng = random.Random(seed)
    chon: List[Dong] = []
    for nhom in NHOM_RUT:
        kho = sorted((r for r in rows if r["phenomenon"] == nhom), key=lambda r: r["uid"])
        if len(kho) < moi_nhom:
            raise ValueError(f"Nhóm {nhom} chỉ có {len(kho)} dòng đã gán, không đủ {moi_nhom}")
        chon.extend(rng.sample(kho, moi_nhom))
    rng.shuffle(chon)
    return chon


def kappa_cohen(a: Sequence[str], b: Sequence[str]) -> Tuple[float, float, float]:
    """Trả về (kappa, p_o, p_e). Hai dãy cùng độ dài, khác rỗng."""
    if len(a) != len(b) or not a:
        raise ValueError("Hai dãy nhãn phải cùng độ dài và khác rỗng")
    n = len(a)
    p_o = sum(x == y for x, y in zip(a, b)) / n
    da, db = Counter(a), Counter(b)
    p_e = sum(da[k] * db[k] for k in set(da) | set(db)) / (n * n)
    if p_e == 1.0:  # cả hai lần chỉ dùng đúng một nhãn: kappa không xác định, quy ước 1 nếu trùng hết
        return (1.0 if p_o == 1.0 else 0.0), p_o, p_e
    return (p_o - p_e) / (1 - p_e), p_o, p_e


def nhan_cam_xuc(r: Dong) -> str:
    """Nhãn cảm xúc mà người chú thích cho là đúng: `dong_y` nghĩa là nhãn gốc."""
    return r["nhan_goc"] if r["nhan_kiem_tra"] == "dong_y" else r["nhan_kiem_tra"]


def doc_xlsx(xlsx: Path) -> List[Dong]:
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True, data_only=True)[ct.TEN_SHEET]
    hang = ws.iter_rows(values_only=True)
    tieu_de = tuple(str(x) for x in next(hang))
    if tieu_de != ct.COT_XLSX:
        raise ValueError(f"Dòng tiêu đề của {xlsx} đã bị đổi: {tieu_de}")
    ra = []
    for h in hang:
        r = {c: ("" if v is None else str(v).strip()) for c, v in zip(ct.COT_XLSX, h)}
        if r["uid"]:
            ra.append(r)
    return ra


def cmd_rut_mau(csv_path: Path, xlsx: Path, moi_nhom: int, seed: int, ghi_de: bool) -> int:
    if xlsx.exists() and not ghi_de:
        print(f"DỪNG: {xlsx} đã tồn tại. Rút lại là đổi mẫu và mất phần đã gán lại. Chắc chắn thì thêm --ghi-de.")
        return 2
    rows = ct.doc_csv(csv_path)
    loi = ct.loi_gia_tri(rows)
    if loi:
        print("CSV có giá trị không hợp lệ, sửa trước:\n  " + "\n  ".join(loi[:10]))
        return 1
    chon = rut_phan_tang(rows, moi_nhom, seed)
    trong = [{**r, **{c: "" for c in ct.COT_CHU_THICH}} for r in chon]  # xoá nhãn cũ khỏi file gán lại
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    ct.luu_workbook(trong, xlsx)
    hom_nay = dt.date.today()
    som_nhat = hom_nay + dt.timedelta(days=SO_NGAY_CHO)
    file_so(xlsx).write_text(json.dumps({
        "ngay_rut": hom_nay.strftime(DINH_DANG_NGAY), "som_nhat_duoc_gan_lai": som_nhat.strftime(DINH_DANG_NGAY),
        "seed": seed, "moi_nhom": moi_nhom, "phan_tang_theo": "phenomenon cua lan gan dau",
        "uid": [r["uid"] for r in chon]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Đã rút {len(chon)} dòng ({moi_nhom} dòng mỗi nhóm {', '.join(NHOM_RUT)}) ra {xlsx}.\n"
          f"Ba cột chú thích trong file này để TRỐNG.\n"
          f"Không mở file trước ngày {som_nhat.strftime(DINH_DANG_NGAY)}. Khi gán lại: không mở candidates.csv, "
          f"không mở candidates_chu_thich.xlsx; chỉ dùng docs/annotation_guideline.md.")
    return 0


def cmd_kappa(csv_path: Path, xlsx: Path, bo_qua_ngay: bool, out_json: Path) -> int:
    so = json.loads(file_so(xlsx).read_text(encoding="utf-8"))
    som_nhat = dt.datetime.strptime(so["som_nhat_duoc_gan_lai"], DINH_DANG_NGAY).date()
    if dt.date.today() < som_nhat and not bo_qua_ngay:
        print(f"DỪNG: mẫu rút ngày {so['ngay_rut']}, sớm nhất {so['som_nhat_duoc_gan_lai']} mới được gán lại "
              f"(cách {SO_NGAY_CHO} ngày để quên nhãn cũ). Nếu người gán là NGƯỜI THỨ HAI thì thêm --bo-qua-ngay.")
        return 2
    lan2 = {r["uid"]: r for r in doc_xlsx(xlsx)}
    lan1 = {r["uid"]: r for r in ct.doc_csv(csv_path)}
    if set(lan2) != set(so["uid"]):
        print("DỪNG: tập uid trong file gán lại không trùng với mẫu đã rút.")
        return 1
    loi = ct.loi_gia_tri(list(lan2.values()))
    thieu = [u for u, r in lan2.items() if not ct.da_xong(r)]
    if loi or thieu:
        print(f"DỪNG: còn {len(thieu)} dòng chưa gán đủ hai ô, {len(loi)} lỗi giá trị.\n  " + "\n  ".join((loi + thieu)[:15]))
        return 1
    uid = so["uid"]
    h1, h2 = [lan1[u]["phenomenon"] for u in uid], [lan2[u]["phenomenon"] for u in uid]
    k, p_o, p_e = kappa_cohen(h1, h2)
    c1, c2 = [nhan_cam_xuc(lan1[u]) for u in uid], [nhan_cam_xuc(lan2[u]) for u in uid]
    kc, pc, _ = kappa_cohen(c1, c2)
    print(f"Hiện tượng (`phenomenon`), {len(uid)} dòng: kappa = {k:.3f} | trùng {p_o * 100:.1f} % | "
          f"trùng do ngẫu nhiên kỳ vọng {p_e * 100:.1f} %")
    print(f"Ngưỡng {NGUONG_KAPPA:.2f}: {'ĐẠT' if k >= NGUONG_KAPPA else 'CHƯA ĐẠT'}")
    nhan = [x for x in ct.HIEN_TUONG if x in set(h1) | set(h2)]
    print("\nLần 1 (hàng) so với lần 2 (cột):")
    print(f"{'':<10}" + "".join(f"{x:>10}" for x in nhan))
    bang = Counter(zip(h1, h2))
    for x in nhan:
        print(f"{x:<10}" + "".join(f"{bang[(x, y)]:>10}" for y in nhan))
    print(f"\nNhãn cảm xúc sau kiểm tra: kappa = {kc:.3f} | trùng {pc * 100:.1f} %")
    lech = [u for u, x, y in zip(uid, h1, h2) if x != y]
    if lech:
        print(f"\n{len(lech)} dòng lệch về hiện tượng (lần 1 -> lần 2):")
        for u in lech:
            print(f"  {u.replace('visfd-test-', '')}: {lan1[u]['phenomenon']} -> {lan2[u]['phenomenon']} | {lan1[u]['text'][:110]}")
    out_json.write_text(json.dumps({
        "loai": "nhat_quan_mot_nguoi" if not bo_qua_ngay else "can_ghi_ro_ai_gan_lan_hai",
        "ngay_rut": so["ngay_rut"], "ngay_tinh": dt.date.today().strftime(DINH_DANG_NGAY), "so_dong": len(uid),
        "phan_tang": True, "moi_nhom": so["moi_nhom"],
        "phenomenon": {"kappa": k, "ti_le_trung": p_o, "ti_le_trung_ngau_nhien": p_e,
                       "bang": {f"{x}->{y}": bang[(x, y)] for x in nhan for y in nhan if bang[(x, y)]}},
        "nhan_cam_xuc": {"kappa": kc, "ti_le_trung": pc}, "uid_lech": lech}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"\nĐã ghi {out_json}")
    return 0 if k >= NGUONG_KAPPA else 3


def kappa_da_dat(kappa_json: Path) -> bool:
    if not kappa_json.exists():
        return False
    return json.loads(kappa_json.read_text(encoding="utf-8"))["phenomenon"]["kappa"] >= NGUONG_KAPPA


def dem_theo_nhom(rows: Sequence[Dong], goc: Path, mo_hinh: Sequence[str]):
    """Trả về (danh sách câu, {mô hình: mảng (nhóm, seed, câu, 9)}).

    Dừng nếu một dòng không có dự đoán, hoặc nhãn thật trong file dự đoán khác `nhan_goc`.
    """
    import numpy as np

    cau = sorted({cb.ma_cau(r["uid"]) for r in rows})
    vi_tri = {c: i for i, c in enumerate(cau)}
    ra = {}
    for exp in mo_hinh:
        o = np.zeros((len(NHOM_CHAM), len(cb.SEEDS), len(cau), cb.SO_LOP * cb.SO_LOP))
        for s, seed in enumerate(cb.SEEDS):
            p = goc / exp / f"seed{seed}" / "predictions.jsonl"
            uid, vang, doan = cb.doc_du_doan(p)
            tra = {u: (int(g), int(d)) for u, g, d in zip(uid, vang, doan)}
            for r in rows:
                if r["uid"] not in tra:
                    raise ValueError(f"{p} không có dự đoán cho {r['uid']}")
                g, d = tra[r["uid"]]
                if g != MA_NHAN[r["nhan_goc"]]:
                    raise ValueError(f"{p}: nhãn thật của {r['uid']} khác nhan_goc trong CSV")
                o[NHOM_CHAM.index(r["phenomenon"]), s, vi_tri[cb.ma_cau(r["uid"])], g * cb.SO_LOP + d] += 1
        ra[exp] = o
    return cau, ra


def cham_mot_mo_hinh(o, so_lan: int, rng, khoi: int = 500, ten_nhom: Sequence[str] = NHOM_CHAM) -> Dict[str, dict]:
    """o: (nhóm, seed, câu, 9). Mỗi nhóm: trung bình ba seed, khoảng tin cậy 95 %, và hiệu so với nhóm mốc (nhóm đầu)."""
    import numpy as np

    so_cau = o.shape[2]
    quan_sat = dict(zip(("macro_f1", "accuracy"), (x.mean(axis=1) for x in cb.do_do(o.sum(axis=2)))))  # (nhóm,)
    mau = {"macro_f1": [], "accuracy": []}
    con = so_lan
    while con > 0:
        n = min(khoi, con)
        w = rng.multinomial(so_cau, np.full(so_cau, 1.0 / so_cau), size=n).astype(np.float64)
        with np.errstate(invalid="ignore", divide="ignore"):  # lần rút không có dòng nào của một nhóm cho nan
            f, a = cb.do_do(np.einsum("nc,gsck->ngsk", w, o))
        mau["macro_f1"].append(f.mean(axis=2))
        mau["accuracy"].append(a.mean(axis=2))
        con -= n
    ra: Dict[str, dict] = {}
    for g, nhom in enumerate(ten_nhom):
        ra[nhom] = {"n": int(o[g, 0].sum())}
        for k in ("macro_f1", "accuracy"):
            m = np.concatenate(mau[k])  # (số lần, nhóm)
            d = m[:, g] - m[:, 0]
            d = d[~np.isnan(d)]
            ra[nhom][k] = {"gia_tri": float(quan_sat[k][g]), "ci95_duoi": float(np.nanpercentile(m[:, g], 2.5)),
                           "ci95_tren": float(np.nanpercentile(m[:, g], 97.5))}
            if g:
                p = 2 * min((np.sum(d <= 0) + 1) / (len(d) + 1), (np.sum(d >= 0) + 1) / (len(d) + 1))
                ra[nhom][k].update(hieu_voi_moc=float(quan_sat[k][g] - quan_sat[k][0]), hieu_ci95_duoi=float(np.percentile(d, 2.5)),
                                   hieu_ci95_tren=float(np.percentile(d, 97.5)), p=float(min(1.0, p)))
    return ra


def _pt(x: float) -> str:
    return f"{x * 100:.1f}".replace(".", ",")


def bang_cham(kq: Dict[str, dict], do_do: str, tieu_de: str, nhom: Sequence[str] = NHOM_CHAM, ten: Dict[str, str] = TEN_NHOM) -> str:
    """Một bảng: hàng là mô hình, cột là nhóm; ô nhóm khác mốc kèm hiệu so với nhóm thường."""
    n = next(iter(kq.values()))
    dong = [tieu_de, "",
            "| Mô hình | " + " | ".join(f"{ten.get(g, g)} (n = {n[g]['n']})" for g in nhom) + " |",
            "|---|" + "---|" * len(nhom)]
    for exp, r in kq.items():
        o = []
        for g in nhom:
            x = r[g][do_do]
            s = f"{_pt(x['gia_tri'])} [{_pt(x['ci95_duoi'])}; {_pt(x['ci95_tren'])}]"
            if "p" in x:
                s += f"; {cb.so(x['hieu_voi_moc'], 1)}, p {cb.p_chu(x['p'])}"
            o.append(s)
        dong.append(f"| {cb.TEN_HIEN_THI.get(exp, exp)} | " + " | ".join(o) + " |")
    return "\n".join(dong) + "\n"


def cmd_cham(csv_path: Path, goc: Path, so_lan: int, seed: int, kappa_json: Path, out_json: Path, out_md: Path, chi_in: bool) -> int:
    import numpy as np

    rows = ct.doc_csv(csv_path)
    loi = ct.loi_gia_tri(rows)
    if loi or not all(ct.da_xong(r) for r in rows):
        print("DỪNG: CSV còn dòng chưa gán hoặc có giá trị không hợp lệ.\n  " + "\n  ".join(loi[:10]))
        return 1
    dung = [r for r in rows if r["phenomenon"] in NHOM_CHAM]
    tap = {"tat_ca": dung, "chi_dong_y": [r for r in dung if r["nhan_kiem_tra"] == "dong_y"]}
    so_bo = not kappa_da_dat(kappa_json)
    rng = np.random.default_rng(seed)
    kq, so_cau = {}, {}
    for ten, ds in tap.items():
        cau, o = dem_theo_nhom(ds, goc, MO_HINH)
        so_cau[ten] = len(cau)
        kq[ten] = {exp: cham_mot_mo_hinh(o[exp], so_lan, rng) for exp in MO_HINH}
    dau = ("**SỐ SƠ BỘ, CHƯA DÙNG CHO BÁO CÁO: chưa có kappa đạt ngưỡng.**\n\n" if so_bo else "")
    chu = ("Mỗi ô: giá trị trung bình ba seed (%), khoảng tin cậy 95 % trong ngoặc vuông; với nhóm khác nhóm thường "
           "có thêm hiệu so với nhóm thường (điểm phần trăm) và p hai phía. Bootstrap {lan} lần theo câu ({cau} câu). "
           "Nhãn chấm là nhãn gốc của bộ dữ liệu.")
    nghin = lambda x: f"{x:,}".replace(",", ".")  # noqa: E731
    khoi = []
    for ten, mo_ta in (("tat_ca", "toàn bộ dòng"), ("chi_dong_y", "chỉ các dòng người chú thích đồng ý với nhãn gốc")):
        for dd, ten_dd in (("accuracy", "Accuracy"), ("macro_f1", "Macro-F1")):
            khoi.append(bang_cham(kq[ten], dd, f"{ten_dd} theo nhóm hiện tượng, {mo_ta}. " + chu.format(lan=nghin(so_lan), cau=nghin(so_cau[ten]))))
    bang = dau + "\n".join(khoi)
    print(bang)
    if not chi_in:
        out_json.write_text(json.dumps({"so_bo": so_bo, "so_lan": so_lan, "seed": seed, "seeds_mo_hinh": list(cb.SEEDS), "don_vi_lay_mau": "cau",
                                        "nhan_cham": "nhan_goc", "nhom_moc": NHOM_CHAM[0], "so_cau": so_cau, "ket_qua": kq},
                                       ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        out_md.write_text(bang, encoding="utf-8")
        print(f"Đã ghi {out_json} và {out_md}")
    return 0


def cmd_chot(csv_path: Path, test_jsonl: Path, out: Path, kappa_json: Path, ghi_de: bool) -> int:
    """Ghi tập chẩn đoán thành .jsonl: mỗi dòng là bản ghi của tập kiểm tra, thêm trường `phenomenon`."""
    if not kappa_da_dat(kappa_json):
        print(f"DỪNG: chưa có {kappa_json} với kappa từ {NGUONG_KAPPA:.2f} trở lên. Chưa đo độ nhất quán thì chưa chốt tập.")
        return 2
    if out.exists() and not ghi_de:
        print(f"DỪNG: {out} đã tồn tại. Chắc chắn muốn ghi lại thì thêm --ghi-de.")
        return 2
    rows = {r["uid"]: r for r in ct.doc_csv(csv_path) if r["phenomenon"] in NHOM_CHAM}
    ra = []
    with test_jsonl.open(encoding="utf-8") as fh:
        for dong in fh:
            b = json.loads(dong)
            r = rows.get(b["uid"])
            if r is None:
                continue
            if b["label"] != MA_NHAN[r["nhan_goc"]]:
                raise ValueError(f"{b['uid']}: nhãn trong {test_jsonl} khác nhan_goc trong CSV")
            ra.append({**b, "phenomenon": r["phenomenon"], "nhan_kiem_tra": r["nhan_kiem_tra"]})
    if len(ra) != len(rows):
        print(f"DỪNG: {len(rows) - len(ra)} dòng của CSV không tìm thấy trong {test_jsonl}.")
        return 1
    out.write_text("".join(json.dumps(b, ensure_ascii=False) + "\n" for b in ra), encoding="utf-8")
    print(f"Đã ghi {out}: {len(ra)} dòng | " + ", ".join(f"{g} {sum(b['phenomenon'] == g for b in ra)}" for g in NHOM_CHAM))
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Rút mẫu gán lại, tính kappa, chấm mô hình và chốt tập chẩn đoán.")
    sub = ap.add_subparsers(dest="lenh", required=True)
    p = sub.add_parser("rut-mau", help="rút mẫu phân tầng, ghi file .xlsx trống nhãn")
    p.add_argument("--csv", type=Path, default=ct.CSV_MAC_DINH)
    p.add_argument("--xlsx", type=Path, default=XLSX_MAC_DINH)
    p.add_argument("--moi-nhom", type=int, default=15, help="số dòng mỗi nhóm hiện tượng (mặc định 15, tức 60 dòng)")
    p.add_argument("--seed", type=int, default=2026)
    p.add_argument("--ghi-de", action="store_true")
    p = sub.add_parser("kappa", help="tính kappa giữa lần gán đầu (CSV) và lần gán lại (.xlsx)")
    p.add_argument("--csv", type=Path, default=ct.CSV_MAC_DINH)
    p.add_argument("--xlsx", type=Path, default=XLSX_MAC_DINH)
    p.add_argument("--out", type=Path, default=Path("data/diagnostic/kappa.json"))
    p.add_argument("--bo-qua-ngay", action="store_true", help="bỏ kiểm tra khoảng cách 14 ngày (khi người gán là người thứ hai)")
    p = sub.add_parser("cham", help="chấm các mô hình theo nhóm hiện tượng từ predictions.jsonl đã có")
    p.add_argument("--csv", type=Path, default=ct.CSV_MAC_DINH)
    p.add_argument("--results", type=Path, default=Path("results"))
    p.add_argument("--so-lan", type=int, default=10_000, help="số lần lấy mẫu bootstrap (mặc định 10000)")
    p.add_argument("--seed", type=int, default=20261013, help="hạt giống của bộ sinh số ngẫu nhiên")
    p.add_argument("--kappa", type=Path, default=KAPPA_MAC_DINH, help="file kappa; thiếu hoặc chưa đạt thì kết quả ghi là sơ bộ")
    p.add_argument("--out-json", type=Path, default=Path("results/diagnostic_eval.json"))
    p.add_argument("--out-md", type=Path, default=Path("chuyende1/tables/bang_4_chan_doan.md"))
    p.add_argument("--chi-in", action="store_true", help="chỉ in bảng, không ghi file")
    p = sub.add_parser("chot", help="ghi file .jsonl của tập chẩn đoán; chỉ chạy khi kappa đã đạt")
    p.add_argument("--csv", type=Path, default=ct.CSV_MAC_DINH)
    p.add_argument("--test", type=Path, default=Path("data/processed/visfd_test.jsonl"))
    p.add_argument("--out", type=Path, default=Path("data/diagnostic/diagnostic_300.jsonl"))
    p.add_argument("--kappa", type=Path, default=KAPPA_MAC_DINH)
    p.add_argument("--ghi-de", action="store_true")
    a = ap.parse_args(argv)
    if a.lenh == "rut-mau":
        return cmd_rut_mau(a.csv, a.xlsx, a.moi_nhom, a.seed, a.ghi_de)
    if a.lenh == "cham":
        return cmd_cham(a.csv, a.results, a.so_lan, a.seed, a.kappa, a.out_json, a.out_md, a.chi_in)
    if a.lenh == "chot":
        return cmd_chot(a.csv, a.test, a.out, a.kappa, a.ghi_de)
    return cmd_kappa(a.csv, a.xlsx, a.bo_qua_ngay, a.out)


if __name__ == "__main__":
    raise SystemExit(main())
